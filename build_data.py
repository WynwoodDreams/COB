#!/usr/bin/env python3
"""Build the opportunity board.

Typical use (from the repo root):

    python3 build_data.py listings.xlsx            # xlsx -> data.json + index.html
    python3 build_data.py --data data.json         # re-render index.html from a saved data.json
    python3 build_data.py listings.xlsx --stats    # also print classification counters
    python3 build_data.py --sheet "Cloud & DevOps" cloud.csv --merge index.html   # add a hand-made sheet

The spreadsheet needs these columns: company, positionName, location,
postingDateParsed, externalApplyLink, salary, description. It can also carry jobType/0..3,
url (the Indeed posting) and id; a row without an id gets one made from its apply link.

A --sheet is a hand-made list of jobs and internships for one spotlight area (SPOT_AREAS),
an .xlsx or a .csv with the columns Company, Role, Location and Direct Apply URL (or the
names in SHEET_ALIASES: Employer, Job title, Indeed application URL...) and optionally
Posted, Type or Job Type, Experience, Salary, a summary column and a caveats column.
Every row lands in that area, and a row that is not an internship is shown as an
entry-level job rather than left out as one.
"""
import argparse, base64, collections, datetime as dt, hashlib, json, os, re, sys

# ---------- category rules: (category, title keywords) in priority order ----------
# A title that names the work decides the area. The Business catch-all sits last and only
# carries words that name business work outright (purchasing, leasing, analyst...). A title
# that matches nothing here is vague, and for those the employer's name and the posting's
# opening paragraph get a say before the catch-all does.
# Spotlight areas: each is fed by a hand-made sheet (--sheet), carries entry-level jobs
# as well as internships, and is drawn to stand out on the page. Their title rules run
# ahead of the Software rule that would otherwise claim "Cloud Engineer", and again in
# finalize(), so cards carried over from older builds move too. The major named here is
# tagged on every row of that area's sheet.
SPOT_AREAS = collections.OrderedDict([
 ("Cloud & DevOps", dict(
    title=r"\bcloud\b|devops|platform engineer|site reliability|\bSRE\b|kubernetes|\bAWS\b|\bazure\b|\bGCP\b",
    major="Information Technology", slug="cloud")),
])
# Areas that used to be spotlights. A card still filed under one is sent back through the
# title classifier in finalize(), so it lands where it would have before the spotlight:
# cyber and help-desk titles in Software, IT, Data & AI, "Cyber Risk Services" in Audit.
RETIRED_AREAS = {"Cybersecurity", "IT Support / Help Desk"}
JOB_LEVEL = "Entry-level job"
CATS = [
 *[(name, spec['title']) for name, spec in SPOT_AREAS.items()],
 ("Health Sciences & Pharmacy", r"pharmac|nurse|nursing|respiratory|medical assistant|\bMA/|chiropract|veterinar|patient care|dietetic|dietitian|nutrition"),
 ("Mental Health & Social Work", r"mental health|counsel|therapist|therapy|psycholog|social work|\bBSW\b|\bMSW\b|RMHCI|LMHC|behavioral health|clinical intern"),
 ("Legal", r"\blaw\b|legal|law clerk|bankruptcy|attorney|paralegal|\b2L\b"),
 ("Accounting, Tax & Audit", r"\btax\b|audit|assurance|accounting|accountant|forensic|attest|risk advisory|technology risk|accounting methods|process risk|internal controls|\bSOX\b"),
 ("Finance, Banking & Insurance", r"financ|wealth|banking|capital markets|private equity|asset management|quant|trading|actuarial|multinational|lending|commercial banking|insurance|state farm agent|FRP internship|client advisory|transaction"),
 ("Human Resources & Talent", r"human resources|\bHR\b|talent acquisition|people & culture|recruit"),
 ("Civil, Structural & Environmental Engineering", r"civil|structural engineer|geotechnical|water|wastewater|transportation|surveying|environmental engineering|\bCEI\b|\bDOT\b|traffic|bridge|structures|intern engineer|engineering & estimating|engineering intern \| fort|planning intern"),
 ("Software, IT, Data & AI", r"software|developer|programmer|\bIT\b|information technology|\bdata\b|analytics|\bAI\b|machine learning|cyber|information security|infosec|security analyst|security engineer|security operations|\bSOC\b|penetration|threat|vulnerabilit|help ?desk|service desk|desktop support|tech(?:nical)? support|support (?:specialist|technician|analyst)|field service tech|system administrator|sys ?admin|digital product|android|web develop|it infrastructure|systems engineering|app / game"),
 ("Architecture, Construction & Design", r"architect|construction|\bBIM\b|interior design|estimating|building automation|fire & life safety|^project intern"),
 ("Mechanical, Electrical & Manufacturing Engineering", r"mechanical|electrical|electronics|\bRF\b|analog|radio|industrial engineering|manufacturing|process engineer|materials|packaging|biomedical|design engineering|engineering intern|engineering internship|engineer, intern|test technician|CAD design|advanced operations|aviation safety|technical sales|(?<!studio )\bengineers?\b"),
 ("Science, Sustainability & Environment", r"chemist|R&D|environmental|sustainab|scientist|research|\blab\b|QEHS|horticult|landscape|grower|nursery|agricultur"),
 ("Marketing, Communications & PR", r"marketing|social media|communication|\bPR\b|public relations|brand|content creator|influencer|campus ambassador|community management|client solutions|e-retail|business communications"),
 ("Creative: Media, Design & Production", r"graphic design|video|photograph|content studio|control room|film|television|studio engineer|music|creative|design intern"),
 ("Hospitality, Events, Sports & Entertainment", r"hospitality|F&B|culinary|hotel|\bclub\b|golf|event|sports|game ops|MIA Academy|strength and conditioning|premium client|sodexo|yotel|academic support|athletic"),
 ("Sales & Business Development", r"\bsales\b|business development|door to door|financial representative"),
 ("Education, Nonprofit & Community", r"ministry|student services|cultivate|programs & partnerships|volunteer|community engagement|community services|nonprofit|school|teaching|teacher|practicum|youth"),
 ("Skilled Trades, Automotive & Aviation", r"mechanic|service technician|aircraft|HVAC|maintenance|trainee.*technician"),
 ("Business, Management & Operations", r"business|operations|supply chain|logistics|purchasing|procurement|consult|administrative|administration|store executive|leasing|customer service|revenue management|analyst|corporate|strategy|leadership development|management trainee|project management|project manager|program development|provider operations"),
]
DEFAULT_CAT = "Business, Management & Operations"
# Titles that say nothing about the work. These skip the title rules entirely.
GENERIC = r"^(intern|internship|intern / co-op|summer internship program.*|internship program.*|.*summer 20\d\d internship)$"
# What the employer is, for a title that doesn't say what the work is.
COMPANY_RULES = [
 ("Architecture, Construction & Design", r"construction|contracting|builders|architecture"),
 ("Legal", r"\blaw\b|legal"),
 ("Science, Sustainability & Environment", r"landscape|nursery|farms?\b"),
 ("Accounting, Tax & Audit", r"advisory|\bCPA"),
 ("Finance, Banking & Insurance", r"\bbank\b|insurance|wealth|financial"),
 ("Hospitality, Events, Sports & Entertainment", r"hotel|resort"),
 ("Health Sciences & Pharmacy", r"hospital|medical center|health system"),
 ("Civil, Structural & Environmental Engineering", r"engineering"),
]
# What the posting's opening says the employer does. Strong nouns only: these run on
# vague titles, where a stray "marketing" would otherwise decide the area.
DESC_HINTS = [
 ("Architecture, Construction & Design", r"general contractor|construction (company|firm|management|site|project|industry|services)|jobsite|job site|ready[- ]mix"),
 ("Legal", r"law (firm|students?|school)|attorneys?|litigation"),
 ("Business, Management & Operations", r"real estate|property management"),
 ("Finance, Banking & Insurance", r"insurance broker|\bbank(ing)?\b|banco|wealth management|investment (firm|management|banking)"),
 # "landscaping" is a trade; "the music landscape" and "the competitive landscape" are not
 ("Science, Sustainability & Environment", r"horticultur|nursery|\bfarms?\b|landscaping|landscape (?:architect|design|contractor|maintenance|management|company|services|industry)"),
 ("Hospitality, Events, Sports & Entertainment", r"\bhotels?\b|\bresorts?\b|hospitality (industry|group|management)"),
 ("Creative: Media, Design & Production", r"record label|music (?:and entertainment )?(?:company|industry|label)"),
]
# Last resort, over more of the posting and with looser words.
DESC_FALLBACK = [
 ("Architecture, Construction & Design", r"construction management|general contractor"),
 ("Accounting, Tax & Audit", r"accounting|\btax\b|audit"),
 ("Software, IT, Data & AI", r"software|computer science|information technology|data analy"),
 ("Marketing, Communications & PR", r"marketing|social media"),
 ("Finance, Banking & Insurance", r"finance|financial"),
 ("Architecture, Construction & Design", r"construction|architect"),
 ("Mechanical, Electrical & Manufacturing Engineering", r"engineering"),
]

# One-off category corrections, for a posting the rules above file somewhere wrong: they
# read a stray word in the description ("audit", "finance", "software"). Each line is
# (employer regex, title regex, area). Add a line here instead of loosening a rule, which
# would move other cards.
CATEGORY_OVERRIDES = [
 (r"^Keith$", r"survey and mapping", "Architecture, Construction & Design"),
 (r"^FTAI Aviation", r"^intern, engineering", "Mechanical, Electrical & Manufacturing Engineering"),
 (r"^StandardAero", r"^student/intern$", "Skilled Trades, Automotive & Aviation"),
 (r"^GA Telesis", r"^summer intern- 2027$", "Business, Management & Operations"),   # one posting for seven departments
 (r"^Southern Glazer", r"\(OneTech\)", "Software, IT, Data & AI"),
 (r"^JM Family", r"^sales planning intern$", "Sales & Business Development"),   # "planning" read as civil
 (r"^Greentech Renewables", r"^intern$", "Sales & Business Development"),      # electrical-distribution sales programme
]

def classify(title, desc, company=''):
    t = title.strip()
    for co, tk, cat in CATEGORY_OVERRIDES:
        if re.search(co, company, re.I) and re.search(tk, t, re.I): return cat
    if re.search(GENERIC, t, re.I):
        for cat, k in COMPANY_RULES:
            if re.search(k, company, re.I): return cat
    for cat, tk in CATS:
        if re.search(tk, t, re.I): return cat
    for cat, k in COMPANY_RULES:
        if re.search(k, company, re.I): return cat
    for cat, k in DESC_HINTS:
        if re.search(k, desc[:1000], re.I): return cat
    for cat, k in DESC_FALLBACK:
        if re.search(k, desc[:1500], re.I): return cat
    return DEFAULT_CAT

# A posting that reached an internship search but is a job: a manager, coordinator or
# supervisor role, or a new-grad hire (HNTB lists those "for former interns only").
NOT_INTERN = re.compile(r"^(manager|director|supervisor)\b|\bnew grad", re.I)
NOT_INTERN_TAIL = re.compile(r"\b(coordinator|manager|director|supervisor)$", re.I)
# Internships as the job's subject matter, not the job: a student services coordinator
# "conducts market outreach, job and internship development".
INTERN_DUTY = re.compile(r"\b(?:jobs?|careers?) and internships?\b|\binternships? (?:development|placements?)\b", re.I)
def is_internship(title, desc):
    if NOT_INTERN.search(title): return False
    # a title that ends in a job word is a job unless the posting itself calls it an internship
    if NOT_INTERN_TAIL.search(display_title(title)) and not re.search(r"\bintern(s|ships?)?\b", INTERN_DUTY.sub(" ", title + " " + desc[:3000]), re.I): return False
    return True

# ---------- majors ----------
# Each pattern is matched against the title and against the degree sentences of the
# posting, one comma-separated phrase at a time, so a lookbehind sees the words before
# the phrase but not the whole posting. (?-i:...) keeps an abbreviation case-sensitive:
# "IT" is a major, "it" is a pronoun.
MAJORS = [
 ("Accounting", r"accounting|accountant|\bCPA\b"),
 ("Finance", r"\bfinance\b|financial planning|asset management|wealth management|investment (banking|management)|actuarial"),
 ("Economics", r"economics"),
 ("Marketing", r"marketing|social media|digital media|advertising"),
 ("Business Administration", r"business administration|business management|business degree|\bMBA\b|management trainee|\bbusiness\b(?![- ](development|communications?|analytics|strategy|needs|transformation|units?|partners?|days?|hours|travel|operations|processes|case|acumen|objectives|goals|results|outcomes|challenges|problems|priorities|value|functions|side|contacts|software|errands|mileage|savvy|owners|activities|information|stakeholders|leaders|forward|intelligence))"),
 ("Communications", r"(?<!verbal )(?<!written )(?<!oral )(?<!strong )(?<!excellent )(?<!good )(?<!effective )(?<!clear )(?<!interpersonal )(?<!professional )(?<!frequent )(?<!client )(?<!guest )communications?\b(?!\s+(skills?|abilities|and (interpersonal|organizational|presentation|collaboration|teamwork)|with|to\b|is\b|are\b|style|tools?|plans?|platforms?|initiatives|channels|team))|journalism|mass comm|communication (studies|arts|sciences)"),
 ("Public Relations", r"public relations"),
 ("Computer Science", r"computer science|software engineer|software develop|computer engineer|\bdeveloper\b"),
 ("Information Technology", r"information technology|information systems|(?-i:\bMIS\b|\bIT\b|\bI\.T\.)(?! (skills|support))"),
 ("Cybersecurity", r"cyber ?security|information security|\bcyber\b"),
 ("Data Science / Analytics", r"data science|data analytics|data engineering|statistics|business analytics|\banalytics\b"),
 ("Artificial Intelligence", r"artificial intelligence|machine learning|(?-i:\bAI\b)(?! tools)"),
 ("Civil Engineering", r"civil engineer|structural engineer|geotechnical|transportation engineer|water resources|\bbridges?\b"),
 ("Mechanical Engineering", r"mechanical (design )?engineer"),
 ("Electrical Engineering", r"electrical engineer|electrical and computer|electronics? engineer"),
 ("Industrial Engineering", r"industrial engineer|manufacturing engineer"),
 ("Biomedical Engineering", r"biomedical(?! informatics)"),
 ("Chemical Engineering", r"chemical engineer"),
 ("Materials Science / Packaging", r"materials (science|engineering)|packaging engineer"),
 ("Environmental Engineering / Science", r"environmental (engineering|science|studies)|marine (biology|science)|oceanograph|\becology\b"),
 ("Aerospace Engineering", r"aerospace|aeronautic"),
 ("Construction Management", r"construction (management|engineering|technology|science|project)|(?-i:\bBIM\b)"),
 ("Architecture", r"(?<!enterprise )(?<!software )(?<!system )(?<!systems )(?<!data )(?<!cloud )(?<!solution )(?<!network )(?<!information )(?<!technical )\barchitect(ure|ural|s)?\b"),
 ("Interior Design", r"interior design"),
 ("Landscape / Horticulture", r"landscape (architecture|design|management)|landscaping|horticultur|agricultur|plant science|\bturf|agronomy|\bgrowers?\b"),
 ("Nursing", r"\bnursing\b|\bnurses?\b|(?-i:\bRN\b|\bBSN\b)"),
 ("Pharmacy", r"pharmacy|pharm\.?d"),
 ("Respiratory Therapy", r"respiratory"),
 ("Health Sciences", r"health science|pre-?med|\bbiology\b|medical assist|chiropract|veterinar|public health|nutrition|dietetic"),
 ("Psychology", r"psychology"),
 ("Counseling", r"counsel(ing|or)|mental health|marriage and family|(?-i:\bLMHC\b|RMHCI)"),
 ("Social Work", r"social work|(?-i:\bMSW\b|\bBSW\b|RCSWI)"),
 ("Graphic Design", r"graphic design|graphic art"),
 ("Film / Video", r"\bfilm\b|\bvideo\b|cinema|broadcast"),
 ("Photography", r"photograph"),
 ("Music", r"\bmusic\b|audio engineering|recording arts|studio engineer"),
 ("Hospitality Management", r"hospitality|hotel management|culinary"),
 ("Sports Management", r"sports management|sport management|kinesiology|exercise science"),
 ("Human Resources", r"human resources?|(?-i:\bHR\b)|people & culture|talent acquisition|organizational (development|leadership)"),
 ("Supply Chain / Logistics", r"supply chain|logistics|procurement|purchasing"),
 ("Mathematics", r"mathematics|\bmath\b"),
 ("Chemistry", r"chemistry"),
 ("Law / Legal Studies", r"law school|(?-i:\bJ\.?D\.?\b)|legal studies|paralegal|pre-?law|law (office|firm|clerk)"),
 ("Education", r"\beducation\b degree|early childhood|teaching|elementary education|secondary education"),
 ("Aviation", r"aviation"),
 ("Real Estate", r"real estate"),
 ("Political Science / Public Admin", r"political science|public administration|public policy"),
]
# Words that open a sentence about what a candidate studies. Bare 'major' is the noisy
# one: unlike the others it is a common adjective ("major market", "three major
# components", and "majority" as a plain substring), so it only counts in the shapes a
# degree requirement actually takes.
_MAJOR_NONACADEMIC = r"market|road|hotel|industr|capital|launch|aspect|component|traffic|citi(?:es)?|compan|brand|metro|platform|social|design"
TRIGGER = re.compile(
    r"degree|majoring|pursuing|enrolled|studying|student in|students in|program in|candidate in|"
    r"background in|coursework|field of study|fields?:|"
    r"law (?:school|students?)|\b[123]L\b|J\.?D\.? candidate|"
    r"bachelor'?s?(?: of \w+)? in|master'?s?(?: of \w+)? in|\bB\.?[AS]\.? in|\bM\.?[AS]\.? in|"
    r"(?:is|are) preferred|preferred in|preferably in|preference for|ideally in|"
    r"(?:undergraduate|graduate|college|university|law|nursing|engineering|pharmacy|doctoral|master'?s|bachelor'?s) students?\b|"
    r"major(?:s|ing)?\s*(?:in\b|:)|"
    r"\b(?:any|all|related|relevant|preferred|desired|recommended|specific|no)\s+majors?\b|"
    r"majors?\s+(?:targeted|represented|include|welcome)|"
    r"major\s+(?:or\s+minor|requirements?)|"
    r"\bmajor(?:s)?\b(?!\s*(?:" + _MAJOR_NONACADEMIC + r"))", re.I)
_BULLET = "•-*·–—"

def _abbrev(desc, i):
    """Is the '.' at desc[i] part of B.S., M.S., Ph.D., e.g.? Then it is not a sentence end."""
    j = i - 1
    while j >= 0 and desc[j].isalpha(): j -= 1
    return 1 <= i - j - 1 <= 2 and (j < 0 or not desc[j].isalnum())

def degree_window(desc, start, end):
    """The sentence a trigger sits in, plus the list it introduces if it ends with one.

    The old rule took 220 characters after the trigger regardless, which reached into
    the next requirement ("Strong communication skills") and the one after that.
    """
    n = len(desc)
    lo = start
    while lo > 0 and start - lo < 100:
        ch = desc[lo - 1]
        if ch == "\n" or (ch in ".;!?" and not _abbrev(desc, lo - 1)): break
        lo -= 1
    hi = end
    while hi < n and hi - end < 340:
        ch = desc[hi]
        if ch in ".?":
            nxt = desc[hi + 1] if hi + 1 < n else " "
            if nxt in " \n" and not (ch == "." and _abbrev(desc, hi)): break
        if ch == ";":
            # "enrolled in a university; preferably in Marketing" runs on, "...; Strong communication" does not
            after = desc[hi + 1:hi + 3].lstrip(" ")
            if not after or after[0] == "\n" or after[0].isupper(): break
        if ch == "\n":
            sofar = desc[lo:hi].rstrip()
            # a header that introduces a list may leave a blank line before it
            header = sofar.endswith(":") or (len(sofar) - (start - lo) < 60 and re.search(r"\bmajors?\b", sofar[start - lo:], re.I))
            rest = desc[hi + 1:].lstrip(" \t")
            if not rest: break
            if rest.startswith("\n"):
                if not header: break
                hi += 1; continue
            line = rest.split("\n", 1)[0].strip()
            # keep going only into a list: "Majors:" then one per line, or bullets, or short items
            if not (header or sofar.endswith(",") or line[:1] in _BULLET or (len(line) <= 45 and "." not in line)): break
        hi += 1
    return lo, hi

_SPLIT = re.compile(r"[,;/•\n()]|\bor\b|\band\b", re.I)
_FILLER = {"a", "an", "the", "or", "and", "in", "of", "to", "with", "from", "e.g", "i.e", "etc", "other", "another",
           "related", "similar", "any", "all", "such", "including", "but", "not", "preferably", "preferred", "ideally"}
KEEP_FIRST = 5   # a posting that names eight degrees "or related" is aimed at the first few
MAX_MAJORS = 6

def majors_for(desc, title):
    """Majors the role is aimed at: the ones in its title, then the first few it names.

    A degree sentence is read one phrase at a time, in order, and only the first
    KEEP_FIRST distinct majors it names are kept. A long "or related field" list is a
    door left open, not a target, and the board does not advertise open doors: a software
    engineering posting that lists Information Security sixth is not a cybersecurity role.
    """
    found = []
    for name, k in MAJORS:
        if re.search(k, title, re.I) and name not in found: found.append(name)
    windows, last = [], -1
    for m in TRIGGER.finditer(desc):
        if m.start() < last: continue          # already inside the previous sentence
        lo, hi = degree_window(desc, m.start(), m.end())
        windows.append(desc[lo:hi]); last = hi
    for w in windows:
        w = re.sub(r"\bJD[/-]MBA\b", "JD", w)   # a JD/MBA candidate is a law student here
        rank, base, i = {}, None, 0
        for tok in _SPLIT.split(w):
            tok = tok.strip()
            if not tok: continue
            hits = [name for name, k in MAJORS if name not in rank and re.search(k, tok, re.I)]
            for name in hits:
                rank[name] = i
                if base is None: base = i
            # a phrase counts as a named degree if it matched one, or reads like one:
            # short, and not a connective ("or a related field", "in a quantitative")
            words = tok.lstrip("*•-–— ").split()
            if hits or (len(words) <= 4 and words and words[0].lower().strip("'’.,") not in _FILLER): i += 1
        for name, i in rank.items():
            if i - base < KEEP_FIRST and name not in found: found.append(name)
    return found[:MAX_MAJORS]

TODAY = dt.date.today()
# Postings advertise the current cycle and the next two years; anything else is a typo or stale.
TERM_YEARS = tuple(str(y) for y in range(TODAY.year, TODAY.year + 3))
# Last month of each academic term, used to tell a term that has already ended from one still ahead.
SEASON_END_MONTH = {"Winter": 2, "Spring": 5, "Summer": 8, "Fall": 12}
def term_is_past(term):
    m = re.match(r"^(Winter|Spring|Summer|Fall)\s+(\d{4})$", term or "")
    if not m: return False
    return (int(m.group(2)), SEASON_END_MONTH[m.group(1)]) < (TODAY.year, TODAY.month)
def term_is_closing(term):
    """A term in its last two months: a 10-15 week internship is no longer recruiting for it."""
    m = re.match(r"^(Winter|Spring|Summer|Fall)\s+(\d{4})$", term or "")
    if not m or term_is_past(term): return False
    y, mo = TODAY.year, TODAY.month + 2
    if mo > 12: y, mo = y + 1, mo - 12
    return (int(m.group(2)), SEASON_END_MONTH[m.group(1)]) <= (y, mo)
def term_for(title, desc):
    pats = r"(Summer|Fall|Spring|Winter|Autumn)\s*(20\d\d)|(20\d\d)\s*(Summer|Fall|Spring|Winter)"
    found = []
    for src in (title, desc[:2500]):
        for m in re.finditer(pats, src, re.I):
            season = (m.group(1) or m.group(4)).title(); year = m.group(2) or m.group(3)
            if season == "Autumn": season = "Fall"
            v = f"{season} {year}"
            if v not in found and year in TERM_YEARS: found.append(v)
        if found: break
    return found[:3]

def clean_loc(l):
    l = s(l)
    l = re.sub(r"\s*\d{5}(-\d{4})?\s*$", "", l).strip()
    return l or "South Florida"

# "Commission" is only pay when the posting talks about earning it. A hospital's
# "Commission on Dietetic Registration" and a casino's "Gaming Commission" are not.
COMMISSION = re.compile(r"commission[- ]based|commission[- ]only|100% commission|plus commission|\+ ?commissions?|"
                        r"earn(?:s|ed|ing)? commissions?|commissions? (?:per|from)\b|uncapped commission|"
                        r"performance[- ]based commission|(?:pay|salary|rate|base) (?:plus|and) commission|commission (?:eligible|structure|pay)", re.I)
JUNK_PAY = re.compile(r"^(up to |from )?\$[0-4](\.\d+)?/(mo|wk|yr|hr)$", re.I)   # "Up to $1/mo" is a placeholder, not pay

def pay_for(sal, desc):
    sal = s(sal)
    if sal:
        sal = sal.replace(" an hour","/hr").replace(" a year","/yr").replace(" a month","/mo").replace(" a week","/wk").replace(" a day","/day")
        return "" if JUNK_PAY.match(sal) else sal
    d = desc[:4000]
    m = re.search(r"\$\s?\d{2}(?:\.\d\d)?\s*(?:-|–|to)\s*\$?\s?\d{2}(?:\.\d\d)?\s*(?:/|per)\s*(?:hour|hr)", d, re.I)
    if m: return m.group(0).replace("per hour","/hr").replace("per hr","/hr").replace(" ","")
    m = re.search(r"\$\s?\d{2}(?:\.\d\d)?\s*(?:/|per)\s*(?:hour|hr)", d, re.I)
    if m: return m.group(0).replace("per hour","/hr").replace("per hr","/hr").replace(" ","")
    if re.search(r"\bpaid internship|\bpaid,|\bpaid\b.{0,20}\bintern|hourly (rate|wage|pay)|stipend|academic credit", d, re.I):
        if re.search(r"unpaid", d, re.I): return "Unpaid / credit"
        return "Paid"
    if re.search(r"\bunpaid\b", d, re.I): return "Unpaid"
    if COMMISSION.search(d): return "Commission"
    return ""

def level_for(title, desc):
    t = title + " " + desc[:3000]
    if re.search(r"law school|\b2L\b|\bJ\.?D\.? candidate|law student", t, re.I): return "Law school"
    if re.search(r"pharm\.?d|doctor of pharmacy|pharmacy (pre-grad )?intern", title + " " + desc[:1500], re.I): return "Pharmacy (PharmD)"
    if re.search(r"registered mental health|registered clinical|RMHCI|LMHC|\bLPC\b|\bMSW\b|school psychologist|counselor intern|therapist intern|Master'?s? Summer Internship|graduate-level|graduate students? only|doctoral|Ph\.?D|supervised clinical|veterinarian", title + " " + desc[:800], re.I): return "Graduate"
    if re.search(r"high school student", t, re.I) and not re.search(r"high school diploma", t, re.I): return "High school"
    return "Undergraduate"

def flag_for(title, desc, pay=""):
    d = title + " " + desc[:5000]
    strong = re.search(r"door[- ]to[- ]door|\b1099\b|no base|commission[- ]only|100% commission|uncapped|make \$", d, re.I)
    if strong or (COMMISSION.search(d) and not re.search(r"/hr", pay or "")): return "Commission-based"
    if re.search(r"for current|returning intern|previous hntb interns", title, re.I): return "Returning interns only"
    return ""

# ---------- languages ----------
# Which languages a posting asks for, and how hard it asks. South Florida postings often ask
# for Spanish, and "a plus" and "required" mean very different things to a student, so the
# wording around each mention is read one clause at a time and the cue nearest the language
# decides. A mention with no cue stays unclassified: the card says "mentioned", not "required".
LANGS = collections.OrderedDict([
 ("Spanish", r"spanish|espa[nñ]ol"),
 ("Portuguese", r"portugu[eê]se?s?\b"),
 ("Haitian Creole", r"creole|krey[oò]l"),
])
BILINGUAL = r"\bbi-?\s?lingual\b|\bbiling[uü]e\b"
OTHER_LANGS = r"mandarin|chinese|cantonese|french|german|russian|arabic|hindi|japanese|korean|vietnamese|italian|hebrew|polish|turkish"
# Boilerplate that names Spanish without asking for it: an E-Verify poster "in English or Spanish",
# a street or a river.
LANG_NOISE = re.compile(r"e-?verify|right to work|poster|notice|click here|equal opportunity|\beeo\b|accommodat|translat|spanish river|spanish (?:trail|wells|springs)", re.I)
LANG_SOFT = re.compile(r"preferred|prefer\b|a plus|\bplus\b|pluses|highly valued|\bvalued\b|bonus|nice to have|desirable|desired|an asset|helpful|ideal|advantage|stand out|not (?:required|mandatory|necessary)|but not|optional|encouraged", re.I)
LANG_HARD = re.compile(r"required|requirement|\bmust\b|mandatory|essential|necessary|needs? to|a must|fully bilingual|(?:seeking|looking for|hiring)\b[^.\n]{0,60}\bbi-?\s?lingual", re.I)
LANG_CLAUSE = re.compile(r"[\n\r\u2022\u00b7*|;]+|(?<=[.!?])\s+|,\s+(?=[A-Za-z(])")
LANG_RANK = {"": 0, "plus": 1, "required": 2}

def langs_for(title, desc):
    """Return ([languages], need) for a posting; need is "required", "plus" or "" (mentioned)."""
    text = (title or "") + "\n" + (desc or "")[:8000]
    found, need = [], ""
    def add(lang, n):
        nonlocal need
        if lang not in found: found.append(lang)
        if LANG_RANK[n] > LANG_RANK[need]: need = n
    named = any(re.search(rx, text, re.I) and not LANG_NOISE.search(text[max(0, m.start()-60):m.end()+60])
                for rx in LANGS.values() for m in re.finditer(rx, text, re.I))
    for i, clause in enumerate(LANG_CLAUSE.split(text)):
        if not clause or LANG_NOISE.search(clause): continue
        hits = [(lang, m) for lang, rx in LANGS.items() for m in [re.search(rx, clause, re.I)] if m]
        if not hits:
            m = re.search(BILINGUAL, clause, re.I)
            # "bilingual" with no language named anywhere is not known to mean Spanish
            if not m or named or re.search(OTHER_LANGS, text, re.I): continue
            hits = [("Bilingual", m)]
        mention = min(m.start() for _, m in hits); end = max(m.end() for _, m in hits)
        def gap(c): return max(c.start() - end, mention - c.end(), 0)
        soft = min((gap(c) for c in LANG_SOFT.finditer(clause)), default=None)
        hard = min((gap(c) for c in LANG_HARD.finditer(clause)), default=None)
        n = ""
        if soft is not None and (hard is None or soft <= hard): n = "plus"
        elif hard is not None: n = "required"
        elif i == 0 and re.search(BILINGUAL, clause, re.I): n = "required"   # "Intern (Bilingual English/Spanish)"
        for lang, _ in hits: add(lang, n)
    return found, need

# Well-known municipalities. These are also stripped from titles, so keep them to names that
# would not plausibly be an ordinary word in a job title.
COUNTY = {
 "Miami-Dade": ["miami","hialeah","coral gables","doral","miami lakes","miami beach","miami gardens","aventura","homestead","north miami","kendall","palmetto bay","cutler bay","opa-locka","medley","key biscayne","sunny isles","pinecrest","south miami","north miami beach","bal harbour","florida city","miami springs","hialeah gardens","sweetwater","miami shores","surfside","west miami"],
 "Broward": ["fort lauderdale","deerfield beach","sunrise","miramar","pompano beach","plantation","hollywood","margate","coral springs","davie","tamarac","weston","hallandale beach","dania beach","cooper city","wilton manors","coconut creek","pembroke pines","lauderhill","oakland park","lighthouse point","parkland","southwest ranches","lauderdale lakes","north lauderdale","pembroke park","lauderdale-by-the-sea"],
 "Palm Beach": ["boca raton","delray beach","west palm beach","boynton beach","jupiter","palm beach gardens","wellington","lake worth","royal palm beach","palm beach","riviera beach","greenacres","north palm beach","juno beach","tequesta"],
}
# Smaller places and unincorporated areas Indeed uses as a location. Only used for the county lookup.
COUNTY_EXTRA = {
 "Miami-Dade": ["gladeview","west little river","westchester","fontainebleau","tamiami","kendale lakes","the hammocks","country walk","richmond west","goulds","princeton","naranja","leisure city","brownsville","golden glades","ives estates","ojus","biscayne park","el portal","virginia gardens","olympia heights","coral terrace","glenvar heights","kendall west","three lakes","south miami heights","west perrine","palmetto estates","the crossings","north bay village","bay harbor islands","indian creek","university park","westwood lakes","sunny isles beach","pinewood","fountainebleau","country club","sunset","miami-dade county","miami-dade"],
 "Broward": ["west park","lazy lake","sea ranch lakes","hillsboro beach","dania","lauderdale by the sea","ft lauderdale","ft. lauderdale","roosevelt gardens","franklin park","washington park","boulevard gardens","broadview park","sunshine ranches","broward county","broward"],
 "Palm Beach": ["lake worth beach","palm springs","lantana","hypoluxo","manalapan","ocean ridge","gulf stream","highland beach","haverhill","loxahatchee","pahokee","belle glade","south bay","lake park","mangonia park","palm beach shores","jupiter inlet colony","atlantis","briny breezes","cloud lake","glen ridge","westlake","palm beach county"],
}
COUNTY_LOOKUP = {city: county for src in (COUNTY, COUNTY_EXTRA) for county, cities in src.items() for city in cities}
def county_for(loc):
    l = re.sub(r",?\s*(fl|florida)\.?$", "", loc.lower().strip()).strip()
    if l in COUNTY_LOOKUP: return COUNTY_LOOKUP[l]
    if "remote" in l: return "Remote"
    if "florida" in l or l in ("fl","south florida"): return "South Florida"
    return "Other Florida"

def norm_company(c):
    c = re.sub(r"\s+-\s+[A-Za-z .]+$", "", c)   # "D1 Training - Coral Springs"
    c = re.sub(r"[.,]", "", c).lower().strip()
    return c
# Indeed sometimes lists one employer under two names, or under the wrong arm of a firm.
# Keys are norm_company() of the name in the export; values are the name the board shows.
# Every name for an employer maps to one, so its postings group and dedupe as one.
COMPANY_NAMES = {
 "team tti": "Techtronic Industries (TTI)",
 "techtronic industries co ltd": "Techtronic Industries (TTI)",
 "baker tilly canada": "Baker Tilly",   # US roles on Baker Tilly's US careers site
}
def company_name(c):
    return COMPANY_NAMES.get(norm_company(c), c)
CITIES = sorted({c for v in COUNTY.values() for c in v} | {"orlando","ft lauderdale","ft. lauderdale","downtown miami","pompano","fort lauderdale","florida","miami-fort lauderdale region"}, key=len, reverse=True)
def norm_title(t):
    t = t.lower().replace("internship","intern")
    for c in CITIES: t = re.sub(r"\b"+re.escape(c)+r"\b", " ", t)
    t = re.sub(r"\b(fl|wm)\b", " ", t)
    t = re.sub(r"[^a-z0-9]+", " ", t).strip()
    return t

def work_mode(loc, desc):
    d = desc[:2500] + " " + s(loc)
    if re.search(r"\bremote\b", d, re.I) and re.search(r"\bhybrid\b", d, re.I): return "Hybrid / Remote"
    if re.search(r"\bhybrid\b", d, re.I): return "Hybrid"
    if re.search(r"fully remote|\(remote\)|remote position|remote intern|work remotely|100% remote", d, re.I): return "Remote"
    return ""

BOILER = r"equal opportunity|EEO|E-Verify|drug[- ]free|background check|third[- ]party agenc|unsolicited resume|benefits? (package|information)|401\(k\)|medical, dental"
SKIP = BOILER + r"|salary range|pay at any point|compensation|pay range|about (us|the company)|is (a|an|the) (leading|global|nation)|we are (a|an|the) (leading|premier|largest)|founded in|headquartered"
def snippet_for(desc):
    lines = [l.strip() for l in desc.split("\n")]
    paras = []
    i = 0
    while i < len(lines):
        p = lines[i]
        if not p: i += 1; continue
        # if a lead-in ends with ":", stitch the following short lines as a clause
        if p.endswith(":") and len(p) > 40:
            tail = []
            j = i + 1
            while j < len(lines) and lines[j] and len(tail) < 3:
                tail.append(lines[j].strip(" •-\t").rstrip(".;")); j += 1
            if tail: p = p[:-1] + ": " + "; ".join(tail) + "."
            i = j
        else:
            i += 1
        if len(p) > 70: paras.append(p)
    pick = None
    for p in paras:
        if re.search(r"\bintern|you will|you'll|responsib|seeking|looking for|opportunity|work (with|alongside)|assist", p, re.I) and not re.search(SKIP, p, re.I) and len(p) > 100:
            pick = p; break
    if not pick:
        for p in paras:
            if not re.search(SKIP, p, re.I): pick = p; break
    if not pick: pick = paras[0] if paras else desc[:240]
    pick = re.sub(r"\s+", " ", pick).strip()
    if len(pick) > 240:
        cut = pick[:240].rsplit(" ", 1)[0]
        pick = cut.rstrip(",;:") + "…"
    return pick

def details_for(desc):
    # grab requirement-ish lines
    lines = [l.strip(" •-\t") for l in desc.split("\n")]
    out = []; grab = False
    for l in lines:
        if re.search(r"^(minimum )?(requirements?|qualifications?|what you.?ll need|what we.?re looking for|who we.?re looking for|what you bring|skills|education|basic qualifications|preferred)", l, re.I) and len(l) < 60:
            grab = True; continue
        if grab:
            if not l:
                if out: break
                continue
            if re.search(BOILER, l, re.I): continue
            if 15 < len(l) < 160: out.append(re.sub(r"\s+"," ",l))
            if len(out) >= 4: break
    return out

CITY_RE = "|".join(re.escape(c) for c in CITIES if c not in ("florida",))
def display_title(t):
    t = re.sub(r"\*+", "", s(t))
    # trailing "- City, FL" / "(City, FL)" / "| City, FL" / "City CLC"
    t = re.sub(r"(?i)\s*[-–|,(]?\s*\b(" + CITY_RE + r")\b(,?\s*FL)?\s*\)?(\s*CLC)?\s*$", "", t)
    # inline "- City, FL -" or "| City, FL |"
    t = re.sub(r"(?i)\s*[-–|]\s*\b(" + CITY_RE + r")\b,?\s*FL\s*([-–|])", r" \2", t)
    # trailing bare state: "... (Summer 2027) FL", "... | Florida", "... - South FL"
    t = re.sub(r"(?i)\s*[-–|,]?\s*\b(south\s+)?(FL|Florida)\b\s*$", "", t)
    t = re.sub(r"\s*\(\s*\)", "", t)
    t = re.sub(r"\s*[-–|,]\s*$", "", t)
    t = re.sub(r"\s{2,}", " ", t).strip(" -–|,")
    while t.endswith(")") and t.count(")") > t.count("("): t = t[:-1].rstrip()
    return t or "Intern"

# ---------- helpers ----------
def s(v):
    """Coerce a spreadsheet cell to a stripped string ('' for None)."""
    if v is None: return ""
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return str(v).strip()

UPGRADED_HTTP = collections.Counter()
def url_or_blank(v, upgrade=True):
    """Keep only http(s) URLs, and prefer the encrypted form.

    A plain-HTTP apply link can be tampered with in transit on shared wi-fi, which is
    exactly the network a student is on. Nearly every employer and ATS serves HTTPS, so
    the upgrade is safe; --keep-http turns it off if a destination ever breaks.
    """
    v = s(v)
    if not re.match(r"^https?://", v, re.I): return ""
    if upgrade and v.lower().startswith("http://"):
        UPGRADED_HTTP[re.sub(r"^http://([^/]+).*", r"\1", v, flags=re.I)] += 1
        v = "https://" + v[len("http://"):]
    return strip_tracking(v)

# Referral tags the scrape picked up on its way through Indeed. They only credit Indeed
# (and in one case carry a per-click session id), and an employer's page opens the same
# without them. Everything else in a query string is left alone: many applicant-tracking
# systems need theirs (opportunityId, jobId, cid...) to find the job at all.
TRACKING_KEYS = re.compile(r"^(utm_\w+|indeed-apply-token|iis|iisn|sid)$", re.I)
SOURCE_KEYS = re.compile(r"^(source|src|rb|lever-source)$", re.I)
def strip_tracking(v):
    base, q, frag = re.match(r"^([^?#]*)(?:\?([^#]*))?(#.*)?$", v).groups()
    if not q: return v
    keep = []
    for part in q.split("&"):
        k, _, val = part.partition("=")
        if TRACKING_KEYS.match(k): continue
        if SOURCE_KEYS.match(k) and re.search(r"indeed|^IN$", val, re.I): continue
        keep.append(part)
    return base + ("?" + "&".join(keep) if keep else "") + (frag or "")

def date_str(v):
    if isinstance(v, (dt.date, dt.datetime)): return v.strftime("%Y-%m-%d")
    v = s(v)[:10]
    return v if re.match(r"^\d{4}-\d{2}-\d{2}$", v) else ""

REQUIRED = ["company", "positionName", "location", "postingDateParsed", "externalApplyLink", "salary", "description"]
# url (the Indeed posting) and id are optional: some exports leave them out. A row without
# either gets an id made from its apply link, so it still de-duplicates and keeps a stable id.

def read_rows(path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("openpyxl is required to read spreadsheets: pip install openpyxl")
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = wb.active.iter_rows(values_only=True)
    try:
        hdr = [s(h) for h in next(rows)]
    except StopIteration:
        sys.exit(f"{path}: spreadsheet is empty")
    missing = [c for c in REQUIRED if c not in hdr]
    if missing: sys.exit(f"{path}: missing columns: {', '.join(missing)}")
    data = []
    for r in rows:
        if not any(s(v) for v in r): continue
        data.append(dict(zip(hdr, r)))
    return data

def made_id(d, company, title):
    key = s(d.get('externalApplyLink')) or "|".join([company, title, s(d.get('location'))])
    return "x" + hashlib.sha1(key.encode("utf-8")).hexdigest()[:12]

NOT_INTERNSHIPS = []
def build_items(data, upgrade=True):
    items, seen, skipped = [], set(), 0
    for d in data:
        title = s(d.get('positionName')); company = company_name(s(d.get('company'))); desc = s(d.get('description'))
        if not title and not company: skipped += 1; continue
        if not is_internship(title, desc): NOT_INTERNSHIPS.append(f"{company} / {title[:50]}"); continue
        pid = s(d.get('id')) or s(d.get('url')) or made_id(d, company, title)
        if pid in seen: skipped += 1; continue
        seen.add(pid)
        indeed = url_or_blank(d.get('url'), upgrade)
        loc = clean_loc(d.get('location'))
        pay = pay_for(d.get('salary'), desc)
        types = []
        for k in ('jobType', 'jobType/0', 'jobType/1', 'jobType/2', 'jobType/3'):
            v = s(d.get(k))
            if v and v not in types: types.append(v)
        items.append(dict(
            id=pid, company=company, title=title or "Intern", loc=loc,
            cat=classify(title, desc, company), majors=majors_for(desc, title), term=term_for(title, desc),
            pay=pay, level=level_for(title, desc), mode=work_mode(d.get('location'), desc),
            types=types, posted=date_str(d.get('postingDateParsed')),
            apply=url_or_blank(d.get('externalApplyLink'), upgrade) or indeed, indeed=indeed,
            snippet=snippet_for(desc), details=details_for(desc), flag=flag_for(title, desc, pay), county=county_for(loc),
        ))
        items[-1]['lang'], items[-1]['langNeed'] = langs_for(title, desc)
    return items, skipped

# ---------- a hand-made sheet for a spotlight area ----------
# Not an Indeed export: one row per role, written by hand for a group that wants one
# kind of role in one place. Jobs are welcome here, so is_internship() is not consulted.
# An .xlsx or a .csv. Each column is found under any of the headers listed for it, so a
# sheet can say "Employer" for Company or "Job title" for Role. Only Company, Role,
# Location and the apply link are required. Posted is optional; a row without one is
# dated the day it is built, which is the day it joined the board. Optional: Type or Job
# Type (an "Internship" is one; Full-time/Part-time is shown, "Full-time graduate program"
# counts as Full-time), Experience (shown on the card), Salary, a Summary (a sentence on
# what the role is, also shown on the card) and Caveats (listed under the card).
SHEET_ALIASES = collections.OrderedDict([
 ("Company", r"^(company|employer)$"),
 ("Role", r"^(role|job title|title|position)$"),
 ("Location", r"^(location|florida location|city)$"),
 ("Direct Apply URL", r"^(direct apply url|apply url|application url|indeed application url|url|link)$"),
 ("Posted", r"^(posted|date posted|posting date|date)$"),
 ("Type", r"^(type|job type|employment|employment type)$"),
 ("Experience", r"^(experience|experience requirement|years)$"),
 ("Salary", r"^(salary|pay|listed pay|compensation)$"),
 ("Summary", r"^(summary|notes?|connection|.*\(from posting\))$"),
 ("Caveats", r"^(caveats?|important caveats|watch out)$"),
])
SHEET_COLUMNS = ["Company", "Role", "Location", "Direct Apply URL"]
NO_PAY = re.compile(r"^(not (listed|stated|given|posted)|n/?a|none|unknown|-+)$", re.I)

def read_sheet_rows(path):
    if re.search(r"\.csv$", path, re.I):
        import csv
        with open(path, encoding='utf-8-sig', newline='') as f:
            raw = list(csv.reader(f))
    else:
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("openpyxl is required to read spreadsheets: pip install openpyxl")
        wb = load_workbook(path, read_only=True, data_only=True)
        raw = list(wb.active.iter_rows(values_only=True))
    if not raw: sys.exit(f"{path}: spreadsheet is empty")
    hdr, rows = [s(h) for h in raw[0]], raw[1:]
    names = {}   # the sheet's header -> the name build_sheet_items() reads
    for canon, pat in SHEET_ALIASES.items():
        for h in hdr:
            if h not in names and re.match(pat, h, re.I): names[h] = canon; break
    missing = [c for c in SHEET_COLUMNS if c not in names.values()]
    if missing: sys.exit(f"{path}: missing columns: {', '.join(missing)} (headers seen: {', '.join(hdr)})")
    out = []
    for r in rows:
        if not any(s(v) for v in r): continue
        d = dict(zip(hdr, r))
        out.append({names.get(h, h): v for h, v in d.items()})
    return out

def build_sheet_items(data, area, upgrade=True):
    spec = SPOT_AREAS[area]
    items, skipped = [], 0
    for d in data:
        title = s(d.get('Role')); company = company_name(s(d.get('Company')).rstrip(" ,;"))
        if not title or not company: skipped += 1; continue
        apply = url_or_blank(d.get('Direct Apply URL'), upgrade)
        if not apply: skipped += 1; continue
        jtype = s(d.get('Type'))
        kind = "Internship" if re.search(r"\bintern", jtype + " " + title, re.I) else "Job"
        shift = re.match(r"(full|part)[- ]time", jtype, re.I)
        exp = s(d.get('Experience'))
        loc = clean_loc(d.get('Location'))
        majors = majors_for("", title)
        if spec['major'] not in majors: majors.insert(0, spec['major'])
        sal = s(d.get('Salary'))
        # The card's chip already says job or internship, so the summary does not repeat it.
        about = []
        if exp: about.append(f"Experience: {exp}")
        if s(d.get('Summary')): about.append(s(d.get('Summary')))
        items.append(dict(
            id=spec['slug'] + "-" + hashlib.sha1(f"{company}|{title}|{apply}".encode('utf-8')).hexdigest()[:10],
            company=company, title=title, loc=loc, cat=area, kind=kind, majors=majors[:MAX_MAJORS],
            term=term_for(title, ""), pay="" if NO_PAY.match(sal) else pay_for(sal, ""),
            level=JOB_LEVEL if kind == "Job" else level_for(title, ""), mode=work_mode(d.get('Location'), ""),
            types=[shift.group(1).title() + "-time"] if shift else [],
            posted=date_str(d.get('Posted')) or TODAY.strftime("%Y-%m-%d"), apply=apply, indeed="",
            snippet=" · ".join(about),
            details=[s(d.get('Caveats'))] if s(d.get('Caveats')) else [], flag="", county=county_for(loc),
        ))
    return items, skipped

# ---------- swapping Indeed links for the employer's own ----------
INDEED_JK = re.compile(r"indeed\.com/(?:viewjob\?(?:[^#]*&)?jk=|job/[^?#]*-)([0-9a-f]{16})\b", re.I)

def hostOf(u):
    m = re.match(r"https?://(?:www\.)?([^/?#]+)", s(u))
    return m.group(1) if m else ""

def indeed_jk(u):
    m = INDEED_JK.search(s(u))
    return m.group(1).lower() if m else ""

def relink_cards(cards, data, upgrade=True):
    """Point cards that apply through Indeed at the employer's own apply page.

    A hand-made sheet often only has the Indeed link. An Indeed export carries the same
    posting's externalApplyLink, matched here by Indeed's job id. The Indeed page stays on
    the card as its "Original listing", and the card takes the export's posting date,
    which is the real one; a sheet row without a date was stamped with the build day.
    """
    ext = {}
    for d in data:
        jk = s(d.get('id')).lower() or indeed_jk(d.get('url'))
        link = url_or_blank(d.get('externalApplyLink'), upgrade)
        if jk and link and not indeed_jk(link) and 'indeed.com' not in link.lower():
            ext[jk] = (link, url_or_blank(d.get('url'), upgrade), date_str(d.get('postingDateParsed')))
    changed = []
    for g in cards:
        hit = False
        for l in g.get('locs') or []:
            jk = indeed_jk(l.get('apply'))
            if jk in ext:
                link, indeed, posted = ext[jk]
                if not s(g.get('indeed')): g['indeed'] = indeed or l['apply']
                l['apply'] = link
                if posted: g['posted'] = posted
                hit = True
        if hit: changed.append(g)
    return changed

def group_items(items):
    """Merge rows that are the same company + title into one card with several locations."""
    groups = collections.OrderedDict()
    for it in items:
        key = (norm_company(it['company']), norm_title(it['title']))
        g = groups.get(key)
        if not g:
            g = dict(it); g['locs'] = [dict(loc=it['loc'], apply=it['apply'], county=it['county'])]; g['count'] = 1; g['counties'] = [it['county']]
            groups[key] = g
        else:
            g['count'] += 1
            if it['loc'] not in [x['loc'] for x in g['locs']]:
                g['locs'].append(dict(loc=it['loc'], apply=it['apply'], county=it['county']))
            if it['county'] not in g['counties']: g['counties'].append(it['county'])
            if not g['flag'] and it['flag']: g['flag'] = it['flag']
            if it['posted'] > g['posted']: g['posted'] = it['posted']
            if not g['pay'] and it['pay']: g['pay'] = it['pay']
            if not g['indeed'] and it['indeed']: g['indeed'] = it['indeed']
            for m in it['majors']:
                if m not in g['majors']: g['majors'].append(m)
            for t in it['term']:
                if t not in g['term']: g['term'].append(t)
            for l in it.get('lang') or []:
                if l not in g.setdefault('lang', []): g['lang'].append(l)
            if LANG_RANK.get(it.get('langNeed') or "", 0) > LANG_RANK.get(g.get('langNeed') or "", 0): g['langNeed'] = it['langNeed']
    out = []
    for g in groups.values():
        for k in ('loc', 'apply', 'county'): g.pop(k, None)
        out.append(g)
    return out

# South Florida industries: the employers this region is known for, found by employer name
# (and, where it helps, title). Each card carries the ids it matches as `hubs`; the board
# filters on them (#hub=<id>) and the insights page shows the collections with at least
# MIN_HUB listings from MIN_HUB_EMP employers, so none can turn into one company's page.
# Check a pattern's matches when you change one: a loose word files strangers under it.
HUBS = [
    dict(id="sports", label="Sports & media", note="Pro teams, stadiums, and Latin music and media companies.",
         co=r"inter miami|miami fc|miami heat|miami dolphins|dolphins stadium|hard rock (stadium|international|live)|grand prix|florida panthers|marlins|miami open|orange bowl|live nation|escol|carbon fiber|fox corporation|fox news|univision|telemundo|sony music|warner music|universal music"),
    dict(id="finance", label="Banking & finance", note="The banks, card companies, wealth managers and insurers of Brickell and beyond.",
         co=r"american express|jpmorgan|j\.?p\.? ?morgan|\bciti\b|citibank|citigroup|wells fargo|bank of america|goldman|morgan stanley|\bubs\b|hsbc|santander|truist|regions bank|city national|bankunited|amerant|raymond james|blackstone|citadel|fidelity|northern trust|mufg|ita[uú]|bbva|banesco|safra|td bank|capital one|\bvisa\b|mastercard|assurant|chubb|\bbank\b|financial|investments?|wealth|asset management|credit union|loop capital"),
    dict(id="realestate", label="Real estate & construction", note="Developers, builders, brokerages and property firms.",
         co=r"related group|lennar|\bmoss\b|coastal construction|suffolk|turner construction|skanska|balfour beatty|kaufman lynn|moriarty|\bancon\b|\bcbre\b|\bjll\b|cushman|colliers(?! engineering)|newmark|berkadia|fortune international|swire|dezer|florida east coast|realty|real estate|propert|\bhomes?\b|builders|construction|dr horton|d\.r\. horton|pulte|toll brothers|gl homes|kolter|minto|codina|greystar|mill creek|avalonbay|hines|tishman|brookfield|prologis|regency centers|kimco"),
    dict(id="aviation", label="Aviation & aerospace", note="Aircraft maintenance, engines, parts, airports and defense electronics.",
         co=r"\bftai\b|ga telesis|heico|standardaero|silver wings|f&e aircraft|l3harris|pratt & whitney|embraer|boeing|lockheed|airlines|jetblue|avianca", ti=r"\baviation\b|aircraft|avionics|aerospace|airport"),
    dict(id="cruise", label="Cruise & shipping", note="Cruise lines and ocean freight.",
         co=r"royal caribbean|carnival|norwegian cruise|windstar|msc cruises|celebrity cruises|seabourn|blue water shipping|seaboard marine|tropical shipping|crowley|port everglades|portmiami", ti=r"\bcruise\b|maritime|shipping"),
]
for _h in HUBS:
    _h["co_re"] = re.compile(_h["co"], re.I); _h["ti_re"] = re.compile(_h["ti"], re.I) if _h.get("ti") else None
MIN_HUB, MIN_HUB_EMP = 5, 3

def hubs_for(g):
    return [h["id"] for h in HUBS if h["co_re"].search(g.get("company", "")) or (h["ti_re"] and h["ti_re"].search(g.get("title", "")))]

def finalize(out, upgrade=True):
    """Idempotent clean-up applied to both fresh and re-loaded data."""
    for g in out:
        g['title'] = display_title(g.get('title'))
        g['indeed'] = url_or_blank(g.get('indeed'), upgrade)
        locs = []
        for l in g.get('locs') or []:
            loc = clean_loc(l.get('loc'))
            locs.append(dict(loc=loc, apply=url_or_blank(l.get('apply'), upgrade) or g['indeed'], county=county_for(loc)))
        g['locs'] = locs
        g['counties'] = list(dict.fromkeys(l['county'] for l in locs))
        g['count'] = max(int(g.get('count') or 1), len(locs))
        for k in ('majors', 'term', 'types', 'details'):
            g[k] = [x for x in (g.get(k) or []) if x]
        for k in ('pay', 'level', 'mode', 'flag', 'snippet', 'company', 'cat', 'kind'):
            g[k] = s(g.get(k))
        g['company'] = company_name(g['company'])
        g['hubs'] = hubs_for(g)
        g['posted'] = date_str(g.get('posted'))
        # Languages: cards from older builds no longer have their full description, so they are
        # read from the title, summary and requirement bullets they do carry. That finds less than
        # the full posting would, and never contradicts it: a fresh build adds to this, never removes.
        langs, need = langs_for(g['title'], "\n".join([g.get('snippet') or ""] + list(g.get('details') or [])))
        have = [l for l in (g.get('lang') or []) if l in (*LANGS, "Bilingual")]
        for l in langs:
            if l not in have: have.append(l)
        if "Bilingual" in have and len(have) > 1: have.remove("Bilingual")   # a named language says it better
        old = g.get('langNeed') if g.get('langNeed') in LANG_RANK else ""
        need = need if LANG_RANK[need] > LANG_RANK[old] else old
        if have: g['lang'] = have
        else: g.pop('lang', None)
        if have and need: g['langNeed'] = need
        else: g.pop('langNeed', None)
        for area, spec in SPOT_AREAS.items():
            if re.search(spec['title'], g['title'], re.I): g['cat'] = area; break
        else:
            # a card left in an area that is no longer a spotlight goes back where its title puts it
            if g['cat'] in RETIRED_AREAS: g['cat'] = classify(g['title'], "", g['company'])
        # Only the spotlight areas mix jobs with internships, so only their cards say which they are.
        if g['cat'] in SPOT_AREAS and not g['kind']: g['kind'] = "Internship"
        # a snippet that only repeats the card's own chip says nothing
        # (cards from older builds also carry it as a prefix, so that comes off too)
        g['snippet'] = re.sub(r"^(?:Entry-level job|Internship)(?: · |$)", "", g['snippet'] or "")
        # An entry-level job keeps its chip wherever it is filed; an internship outside a
        # spotlight area is the default and needs none.
        if g['cat'] not in SPOT_AREAS and g['kind'] != "Job": g.pop('kind', None)
    out.sort(key=lambda x: (x['company'].lower(), x['title'].lower()))
    out.sort(key=lambda x: x['posted'], reverse=True)   # stable: newest first, then employer A-Z
    return out

# Postings kept off the board on purpose. Each line is (employer regex, title regex, why).
# Held here, not just deleted from index.html, so the next export cannot bring them back.
EXCLUDED = [
 (r"^Lotus Sales", r".", "door-to-door commission sales sold as an internship"),
 (r"^Hongfeng US Services", r"office intern", "the job is in Irvine, California; Indeed filed it under Miami"),
 (r"^Planate Management", r"robotics intern", "the job is in Orlando; Indeed filed it under Miami"),
 (r"^king cable", r"intern/office assistant", "worded as a permanent office job, not an internship"),
 (r"^Saludablemente", r"mental health counselor", "hires licensed counselors; not a student internship"),
 (r"^RI Creative Studio", r"content creator manager/intern", "a manager-style title posted by one person hiring for their own team"),
 (r"^Baker Tilly", r"business valuation associate", "a salaried full-time associate role, not an internship"),
 (r"^Oasis Futsal", r"futsal coach", "a part-time coaching job, not an internship"),
 (r"^TBD Influencer", r"personal assistant intern", "unpaid personal-assistant work for an unnamed individual"),
]

def drop_excluded(out):
    kept, held = [], []
    for g in out:
        why = next((w for co, tk, w in EXCLUDED if re.search(co, g['company'], re.I) and re.search(tk, g['title'], re.I)), None)
        (held if why else kept).append((g, why) if why else g)
    return kept, held

def drop_expired(out):
    """Strip terms that have already ended, and drop a card whose only term has ended.

    A card that never named a term is left alone: it makes no claim about a cycle.
    """
    kept, expired = [], []
    for g in out:
        live = [t for t in g.get('term') or [] if not term_is_past(t)]
        past = [t for t in g.get('term') or [] if term_is_past(t)]
        if past and not live:
            expired.append(g); continue
        g['term'] = live
        kept.append(g)
    return kept, expired

def drop_closing(out, confirmed):
    """Drop a card whose every term is closing unless the fresh export still lists it.

    Nothing in an export says a posting closed, but a Fall card the October scrape no
    longer carries, for a term that ends in December, is far more likely gone than open.
    A card the export did confirm, or one that also names a later term, stays.
    """
    kept, closing = [], []
    for g in out:
        terms = g.get('term') or []
        if terms and all(term_is_closing(t) for t in terms) and g.get('id') not in confirmed:
            closing.append(g)
        else:
            kept.append(g)
    return kept, closing

def drop_stale(out, cutoff):
    """Split cards by whether their newest posting predates the cutoff. Dead links cost trust."""
    kept, stale = [], []
    for g in out:
        (stale if g.get('posted') and g['posted'] < cutoff else kept).append(g)
    return kept, stale

# ---------- merging a fresh export into the board already published ----------
def load_cards(path):
    """Read cards from a data.json, or from the inline JSON of a built index.html.

    data.json is gitignored, so a committed index.html is usually the only record of
    what is on the board; being able to read it back is what makes --merge additive.
    """
    text = open(path, encoding='utf-8').read()
    if re.search(r"\.html?$", path, re.I):
        m = re.search(r'<script id="data" type="application/json">(.*?)</script>', text, re.S)
        if not m: sys.exit(f"{path}: no <script id=\"data\"> block to read listings from")
        text = m.group(1).replace("\\u003c", "<")
    try:
        cards = json.loads(text)
    except json.JSONDecodeError as e:
        sys.exit(f"{path}: not valid JSON ({e})")
    if not isinstance(cards, list): sys.exit(f"{path}: expected a JSON list of cards")
    return cards

def card_key(g):
    """Same identity a build uses to group rows: employer + title, city and case stripped."""
    return (norm_company(s(g.get('company'))), norm_title(display_title(g.get('title'))))

def card_alias(g):
    """Looser key: the same words in any order, for a role reposted with a shuffled title.

    "Intern" and the words of the card's own term are left out of the title, and the terms
    join the key instead. "TTI Group Finance and Accounting Internship" meets the same title
    with "- Summer 2027" appended when both are for Summer 2027, while "Audit Summer 2027"
    and "Audit Winter 2027" stay two roles. A season the build could not read as a term
    ("Audit Spring Internship 2027") stays in the title words, so it still tells roles apart.
    """
    co, t = card_key(g)
    terms = frozenset(g.get('term') or [])
    said = {w.lower() for x in terms for w in x.split()} | {"intern", "interns"}
    return (co, frozenset(w for w in t.split() if w not in said), terms)

def merge_into(g, n):
    """Fold a fresh card into the one already on the board, and say whether anything changed.

    The published card wins on prose, because it is the one that has been eyeballed; lists
    take the union, and a newer posting date wins so a refreshed listing floats back up.
    """
    before = json.dumps(g, sort_keys=True, ensure_ascii=False)
    have = [l.get('loc') for l in g.get('locs') or []]
    for l in n.get('locs') or []:
        if l.get('loc') not in have:
            g.setdefault('locs', []).append(dict(l)); have.append(l.get('loc'))
    for k in ('term', 'majors', 'types', 'counties'):
        vals = list(g.get(k) or [])
        for v in n.get(k) or []:
            if v not in vals: vals.append(v)
        g[k] = vals
    for k in ('pay', 'flag', 'indeed', 'mode', 'snippet', 'level', 'cat', 'kind'):
        if not s(g.get(k)) and s(n.get(k)): g[k] = n[k]
    # A hand-made sheet is curated, so its say on the area beats a classifier's guess.
    if s(n.get('cat')) in SPOT_AREAS and s(n.get('kind')):
        g['cat'] = n['cat']; g['kind'] = n['kind']
    if not (g.get('details') or []) and n.get('details'): g['details'] = list(n['details'])
    for l in n.get('lang') or []:
        if l not in g.setdefault('lang', []): g['lang'].append(l)
    if LANG_RANK.get(n.get('langNeed') or "", 0) > LANG_RANK.get(g.get('langNeed') or "", 0): g['langNeed'] = n['langNeed']
    if s(n.get('posted')) > s(g.get('posted')): g['posted'] = n['posted']
    g['count'] = max(int(g.get('count') or 1), int(n.get('count') or 1), len(g.get('locs') or []))
    return json.dumps(g, sort_keys=True, ensure_ascii=False) != before

def merge_cards(prior, fresh):
    """Add only what the board does not already carry.

    A fresh card matches a published one by employer + title, by Indeed posting id, or by
    the same title words in another order. A match updates that card instead of adding a
    second one, so re-running a build over an overlapping export cannot duplicate a role.
    """
    out = [dict(g) for g in prior]
    by_key, by_alias, by_ref = {}, {}, {}   # by_ref holds posting ids and apply links alike
    def index(g):
        by_key.setdefault(card_key(g), g); by_alias.setdefault(card_alias(g), g)
        for ref in [g.get('id')] + [l.get('apply') for l in g.get('locs') or []]:
            if s(ref): by_ref.setdefault(s(ref), g)
    for g in out: index(g)
    added, updated, same = [], [], []
    for n in fresh:
        refs = [n.get('id')] + [l.get('apply') for l in n.get('locs') or []]
        g = by_key.get(card_key(n)) or by_alias.get(card_alias(n))
        for ref in refs:
            if g is not None: break
            cand = by_ref.get(s(ref)) if s(ref) else None
            # one careers page can serve several roles at an employer, so a shared apply
            # link only settles it when the employer agrees too
            if cand is not None and norm_company(s(cand.get('company'))) == norm_company(s(n.get('company'))):
                g = cand
        if g is None:
            c = dict(n); out.append(c); index(c); added.append(c)
        elif merge_into(g, n): updated.append(g)
        else: same.append(g)
    return out, added, updated, same

def dedupe_cards(cards):
    """Fold cards that are one role into the newest of them, across the whole board.

    merge_cards only compares fresh cards with published ones, so two copies that are both
    already published, or both in one export under different titles, would otherwise stay.
    Matching is by employer + title or by card_alias; a shared apply link is not enough on
    its own here, because one careers page can list every role an employer has.
    """
    kept, folded, seen = [], [], {}
    for g in cards:   # newest first, from finalize, so the newest copy is the one kept
        k = seen.get(card_key(g)) or seen.get(card_alias(g))
        if k is None:
            kept.append(g); seen.setdefault(card_key(g), g); seen.setdefault(card_alias(g), g)
        else:
            merge_into(k, g); folded.append(g)
    return kept, folded

# ---------- featured (paid) placements ----------
# featured.json, next to this script, is the list of roles an employer has paid to pin.
# Each entry names the card by company and title (matched the way merge_cards matches) and
# the last day it stays featured. An entry that also gives an apply link puts the role on
# the board even when no export carries it, so a sponsor's role need not wait for a scrape.
# The file is the whole truth: every build clears the mark and sets it again from here.
FEATURE_FIELDS = {"company", "title", "until", "apply", "location", "pay", "summary", "term"}

def load_featured(path):
    if not os.path.exists(path): return []
    try:
        entries = json.load(open(path, encoding='utf-8'))
    except json.JSONDecodeError as e:
        sys.exit(f"{path}: not valid JSON ({e})")
    if not isinstance(entries, list): sys.exit(f"{path}: expected a JSON list of entries")
    for i, e in enumerate(entries):
        where = f"{path} entry {i+1}"
        if not isinstance(e, dict): sys.exit(f"{where}: expected an object")
        extra = set(e) - FEATURE_FIELDS
        if extra: sys.exit(f"{where}: unknown field(s) {', '.join(sorted(extra))}; allowed: {', '.join(sorted(FEATURE_FIELDS))}")
        if not s(e.get('company')) or not s(e.get('title')): sys.exit(f"{where}: needs a company and a title")
        if not date_str(e.get('until')): sys.exit(f"{where}: 'until' must be a date like 2026-11-30")
    return entries

def apply_featured(cards, entries, upgrade=True):
    """Mark the cards named in featured.json, adding any that carry their own apply link.

    Returns (featured, added, missing, ended) for the build log.
    """
    for g in cards: g.pop('featured', None)
    today = TODAY.strftime("%Y-%m-%d")
    by_key = {}
    for g in cards:
        by_key.setdefault(card_key(g), g); by_key.setdefault(card_alias(g), g)
    featured, added, missing, ended = [], [], [], []
    for e in entries:
        until = date_str(e['until'])
        probe = dict(company=company_name(s(e['company'])), title=display_title(s(e['title'])), term=list(e.get('term') or []))
        g = by_key.get(card_key(probe)) or by_key.get(card_alias(probe))
        if until < today:
            ended.append(e); continue
        if g is None:
            apply = url_or_blank(e.get('apply'), upgrade)
            if not apply:
                missing.append(e); continue
            loc = clean_loc(e.get('location')) or "South Florida"
            summary = s(e.get('summary'))
            g = dict(
                id="feat-" + hashlib.sha1(f"{probe['company']}|{probe['title']}|{apply}".encode('utf-8')).hexdigest()[:10],
                company=probe['company'], title=probe['title'], cat=classify(probe['title'], summary, probe['company']),
                majors=majors_for(summary, probe['title']), term=probe['term'] or term_for(probe['title'], summary),
                pay=pay_for(e.get('pay'), "") if s(e.get('pay')) else "", level=level_for(probe['title'], summary),
                mode=work_mode(e.get('location'), summary), types=[], posted=today, indeed="",
                snippet=summary, details=[], flag="", count=1, counties=[county_for(loc)],
                locs=[dict(loc=loc, apply=apply, county=county_for(loc))],
            )
            if not is_internship(probe['title'], summary): g['kind'] = "Job"; g['level'] = JOB_LEVEL
            cards.append(g); by_key.setdefault(card_key(g), g); by_key.setdefault(card_alias(g), g)
            added.append(g)
        g['featured'] = until
        featured.append(g)
    return featured, added, missing, ended

# Descriptions are written by whoever posted the job, and they land in this repo where
# coding agents read them. None of it reaches an LLM at runtime, so this is an early
# warning for the humans and agents who work on the repo, not a runtime defence.
SUSPECT = re.compile(
    r"ignore (all |any )?(previous|prior|above)\b|disregard (the |all )?(above|previous)|"
    r"\bsystem prompt\b|you are now (a|an|in)\b|new instructions?:|</?system>|\[\[?INST\]?\]|"
    r"^\s*(assistant|system)\s*:|\bexfiltrat|\bapi[ _-]?key\b|\bbase64\s+-d\b|"
    r"\beval\(|\bprocess\.env\b|curl\s+https?://", re.I | re.M)
SCAN_FIELDS = ("company", "title", "snippet", "details", "majors", "term", "pay", "flag", "level")

def scan_injection(cards):
    """Flag listing text shaped like an instruction aimed at an AI agent."""
    hits = []
    for g in cards:
        for f in SCAN_FIELDS:
            v = g.get(f)
            for text in (v if isinstance(v, list) else [v]):
                m = SUSPECT.search(s(text))
                if m: hits.append((g.get("company", ""), g.get("title", ""), f, m.group(0).strip()))
    return hits


# ---------- the insights page ----------
# Skills the insights page counts, by scanning each card's title, summary and requirements
# once (a card mentions a skill or it does not, however many times it says so). Short and
# curated on purpose: a free-text keyword count fills up with "team" and "ability".
SKILLS = [
    ("Excel", r"\bexcel\b"), ("PowerPoint", r"\bpowerpoint\b"), ("Microsoft Office", r"\b(microsoft|ms) office\b|\boffice suite\b|\bmicrosoft 365\b"),
    ("Communication skills", r"\bcommunication skills?\b|\bcommunicat(e|ion) (clearly|effectively)\b|\b(written|verbal) communication"),
    ("Data analysis", r"\bdata anal(ysis|ytics)\b|\banalyz(e|ing) data\b"), ("Project management", r"\bproject management\b"),
    ("Customer service", r"\bcustomer service\b"), ("Social media", r"\bsocial media\b"), ("Research", r"\bresearch\b"),
    ("Python", r"\bpython\b"), ("SQL", r"\bsql\b"), ("Java", r"\bjava\b"), ("JavaScript", r"\bjavascript\b|\btypescript\b|\breact\b|\bnode\.?js\b"),
    ("C / C++ / C#", r"\bc\+\+|\bc#|\bc programming\b"), ("R", r"\bR\b(?= |,|\.|/)(?=.*\b(statistic|stata|spss|data)\b)"),
    ("Tableau / Power BI", r"\btableau\b|\bpower ?bi\b"), ("AWS / Azure / GCP", r"\baws\b|\bazure\b|\bgoogle cloud\b|\bgcp\b"),
    ("Linux", r"\blinux\b"), ("Git / GitHub", r"\bgit(hub)?\b"), ("Machine learning / AI", r"\bmachine learning\b|\bartificial intelligence\b|\bgenerative ai\b|\bllm"),
    ("Salesforce", r"\bsalesforce\b"), ("SAP", r"\bsap\b"), ("QuickBooks", r"\bquickbooks\b"), ("Bloomberg", r"\bbloomberg\b"),
    ("AutoCAD", r"\bautocad\b"), ("Revit", r"\brevit\b"), ("SolidWorks", r"\bsolidworks\b"), ("MATLAB", r"\bmatlab\b"), ("GIS", r"\barcgis\b|\bgis\b"),
    ("Adobe Creative Cloud", r"\badobe\b|\bphotoshop\b|\billustrator\b|\bindesign\b|\bpremiere\b|\bafter effects\b"), ("Canva", r"\bcanva\b"), ("Figma", r"\bfigma\b"),
    ("Spanish", r"\bspanish\b"), ("Bilingual", r"\bbilingual\b"),
    ("GPA requirement", r"\bgpa\b"), ("Driver's license", r"\bdriver'?s? licen[sc]e\b|\bvalid licen[sc]e\b"),
    ("Writing skills", r"\bwriting skills?\b|\bstrong writ(er|ing)\b|\bcopywriting\b"), ("Public speaking", r"\bpublic speaking\b|\bpresentation skills?\b"),
    ("Lab experience", r"\blaborator(y|ies)\b|\bwet lab\b|\bbench work\b"), ("CPR / BLS", r"\bcpr\b|\bbls\b"),
]
SKILL_RE = [(label, re.compile(rx, re.I)) for label, rx in SKILLS]
MIN_PAY_N = 5          # same bar as the pay report in template.html
MIN_SKILL_N = 2
TOP_EMPLOYERS = 12

def hourly_quote(g):
    """The hourly pay a card quotes, for the pay figures: the middle of a range, hourly only,
    internships only, mirroring hourlyQuote in template.html."""
    pay = g.get("pay") or ""
    if g.get("kind") == "Job" or not re.search(r"/hr\b", pay, re.I): return None
    nums = [float(x.replace(",", "")) for x in re.findall(r"\d[\d,]*\.?\d*", pay)]
    nums = [x for x in nums if x == x]
    if not nums: return None
    v = (min(nums) + max(nums)) / 2
    return v if 8 <= v <= 120 else None

def quant(vals, q):
    a = sorted(vals); i = (len(a) - 1) * q; lo = int(i // 1); hi = min(lo + 1, len(a) - 1)
    return a[lo] + (a[hi] - a[lo]) * (i - lo)

def pay_stats(cards):
    v = [x for x in (hourly_quote(g) for g in cards) if x is not None]
    return {"n": len(v), "med": round(quant(v, .5), 2), "lo": round(quant(v, .25), 2), "hi": round(quant(v, .75), 2)} if v else None

def pay_rows(groups, key):
    """Pay stats per group, the groups with too few hourly figures listed by name instead."""
    rows, thin = [], []
    for label, cards in groups.items():
        st = pay_stats(cards)
        if not st: continue
        (rows if st["n"] >= MIN_PAY_N else thin).append(dict({key: label}, **st))
    rows.sort(key=lambda r: (-r["med"], r[key]))
    return rows, sorted(r[key] for r in thin)

def card_text(g):
    return " ".join([g.get("title", ""), g.get("snippet", "")] + list(g.get("details") or []))

def skills_in(cards):
    c = collections.Counter()
    for g in cards:
        t = card_text(g)
        c.update(label for label, rx in SKILL_RE if rx.search(t))
    return [[k, n] for k, n in c.most_common() if n >= MIN_SKILL_N]

def days_since(d):
    try: return (TODAY - dt.date.fromisoformat(d)).days
    except (TypeError, ValueError): return None

def insights_for(cards):
    """Everything the insights page shows, worked out once here so the page itself stays
    small: it carries these figures, not the listings."""
    interns = [g for g in cards if g.get("kind") != "Job"]
    ages = [a for a in (days_since(g.get("posted")) for g in cards) if a is not None]
    new14 = sum(a <= 14 for a in ages); prev14 = sum(14 < a <= 28 for a in ages)
    by = lambda f: collections.Counter(f(g) for g in cards if f(g))
    multi = lambda f: collections.Counter(v for g in cards for v in (f(g) or []))
    counties = multi(lambda g: g.get("counties"))
    terms = multi(lambda g: g.get("term"))
    levels = by(lambda g: g.get("level"))
    areas = by(lambda g: g.get("cat"))
    paid = [g for g in cards if g.get("pay")]
    pay_all = pay_stats(interns)
    pulse = {
        "open": len(cards), "interns": len(interns), "jobs": len(cards) - len(interns),
        "openings": sum(int(g.get("count") or 1) for g in cards),
        "employers": len({g["company"] for g in cards}), "areas": len(areas),
        "new7": sum(a <= 7 for a in ages), "new14": new14, "prev14": prev14, "new30": sum(a <= 30 for a in ages),
        "payListed": len(paid), "hourlyN": pay_all["n"] if pay_all else 0, "medHr": pay_all["med"] if pay_all else None,
        "remote": sum(1 for g in cards if g.get("mode")), "bilingual": sum(1 for g in cards if g.get("lang")),
        "counties": counties.most_common(), "terms": terms.most_common(), "levels": levels.most_common(),
        "areas": areas.most_common(),
    }
    group = lambda f: {k: [g for g in interns if k in (f(g) or [])] for k in {v for g in interns for v in (f(g) or [])}}
    area_rows, area_thin = pay_rows({a: [g for g in interns if g.get("cat") == a] for a in areas}, "area")
    county_rows, _ = pay_rows(group(lambda g: g.get("counties")), "county")
    level_rows, _ = pay_rows({l: [g for g in interns if g.get("level") == l] for l in levels}, "level")
    emp_rows, _ = pay_rows({c: [g for g in interns if g["company"] == c] for c in {g["company"] for g in interns}}, "employer")
    pay = {"all": pay_all, "areas": area_rows, "thin": area_thin, "counties": county_rows, "levels": level_rows, "employers": emp_rows[:10]}

    employers = []
    for name, n in collections.Counter(g["company"] for g in cards).most_common(TOP_EMPLOYERS):
        cs = [g for g in cards if g["company"] == name]
        rates = [x for x in (hourly_quote(g) for g in cs) if x is not None]
        cats = collections.Counter(g.get("cat") for g in cs).most_common()
        newest = min((a for a in (days_since(g.get("posted")) for g in cs) if a is not None), default=None)
        employers.append({
            "name": name, "cards": n, "openings": sum(int(g.get("count") or 1) for g in cs),
            "areas": [c for c, _ in cats], "counties": [c for c, _ in collections.Counter(c for g in cs for c in g.get("counties") or []).most_common()],
            "payLo": round(min(rates)) if rates else None, "payHi": round(max(rates)) if rates else None,
            "paid": sum(1 for g in cs if g.get("pay") and not re.search(r"unpaid", g["pay"], re.I)),
            "bilingual": any(g.get("lang") for g in cs), "remote": any(g.get("mode") for g in cs),
            "jobs": sum(1 for g in cs if g.get("kind") == "Job"), "newest": newest,
            "terms": [t for t, _ in collections.Counter(t for g in cs for t in g.get("term") or []).most_common(2)],
        })
    hubs = []
    for h in HUBS:
        cs = [g for g in cards if h["id"] in (g.get("hubs") or [])]
        emps = collections.Counter(g["company"] for g in cs)
        if len(cs) < MIN_HUB or len(emps) < MIN_HUB_EMP: continue
        st = pay_stats([g for g in cs if g.get("kind") != "Job"])
        hubs.append({"id": h["id"], "label": h["label"], "note": h["note"], "n": len(cs), "employers": len(emps),
                     "openings": sum(int(g.get("count") or 1) for g in cs), "medHr": st["med"] if st and st["n"] >= MIN_PAY_N else None, "hourlyN": st["n"] if st else 0,
                     "areas": [a for a, _ in collections.Counter(g.get("cat") for g in cs).most_common(3)],
                     "topEmployers": [[c, n] for c, n in emps.most_common(3)], "bilingual": sum(1 for g in cs if g.get("lang")),
                     "paid": sum(1 for g in cs if g.get("pay") and not re.search(r"unpaid", g["pay"], re.I))})
    skills = {"all": skills_in(cards), "byArea": {a: skills_in([g for g in cards if g.get("cat") == a])[:8] for a in areas}, "scanned": len(cards)}
    return {"pulse": pulse, "pay": pay, "employers": employers, "hubs": hubs, "skills": skills}

# ---------- history: one snapshot of the figures per build, kept in the repo ----------
# history.json is the board's record over time: every build adds (or, on the same day,
# replaces) one line of headline numbers, so quarters and years can be compared later. It
# is committed with the pages, so git keeps every version and nothing needs a server.
def snapshot_for(cards, ins):
    P, pay = ins["pulse"], ins["pay"]
    keep = ("open", "interns", "jobs", "openings", "employers", "new7", "new14", "new30",
            "payListed", "hourlyN", "medHr", "remote", "bilingual")
    snap = {"date": TODAY.isoformat()}
    snap.update({k: P[k] for k in keep})
    snap.update({
        "payLo": pay["all"]["lo"] if pay["all"] else None, "payHi": pay["all"]["hi"] if pay["all"] else None,
        "areas": dict(P["areas"]), "counties": dict(P["counties"]), "terms": dict(P["terms"]), "levels": dict(P["levels"]),
        "payAreas": {r["area"]: r["med"] for r in pay["areas"]},
        "topEmployers": [[e["name"], e["cards"]] for e in ins["employers"][:5]],
        "skills": ins["skills"]["all"][:10],
        "hubs": {h["id"]: h["n"] for h in ins["hubs"]},
        # postings by the month they were posted, so a quarter's total survives the 90-day trim
        "postedByMonth": dict(collections.Counter(g["posted"][:7] for g in cards if g.get("posted"))),
    })
    return snap

def update_history(path, snap):
    hist = json.load(open(path, encoding='utf-8')) if os.path.exists(path) else []
    hist = sorted([h for h in hist if h["date"] != snap["date"]] + [snap], key=lambda h: h["date"])
    with open(path, 'w', encoding='utf-8') as f:   # one snapshot per line, so diffs read
        f.write("[\n" + ",\n".join(json.dumps(h, ensure_ascii=False, separators=(",", ":")) for h in hist) + "\n]\n")
    return hist

def quarter_of(d):
    return f"{d[:4]} Q{(int(d[5:7]) - 1) // 3 + 1}"

def load_visits(path):
    """visits.json, written by record_visits.py from Vercel Web Analytics: page views per day
    and page. Missing or empty means the weekly job has not run yet."""
    if not path or path == '-' or not os.path.exists(path): return {}
    try: return json.load(open(path, encoding='utf-8')).get("days") or {}
    except (json.JSONDecodeError, AttributeError): return {}

def visits_series(days):
    """Views by month and by week (all pages, and the insights page on its own)."""
    if not days: return None
    by_month, by_week = {}, {}
    for d, v in sorted(days.items()):
        m = d[:7]
        wk = (dt.date.fromisoformat(d) - dt.timedelta(days=dt.date.fromisoformat(d).weekday())).isoformat()  # that week's Monday
        for key, bucket in ((m, by_month), (wk, by_week)):
            b = bucket.setdefault(key, {"all": 0, "insights": 0})
            b["all"] += int(v.get("board", 0)) + int(v.get("insights", 0)) + int(v.get("other", 0))
            b["insights"] += int(v.get("insights", 0))
    return {"months": [[m, b["all"], b["insights"]] for m, b in sorted(by_month.items())][-12:],
            "weeks": [[w, b["all"], b["insights"]] for w, b in sorted(by_week.items())][-16:],
            "first": min(days), "last": max(days), "days": len(days)}

def history_series(hist, visits=None):
    """What the insights page draws from the history: a point per build over the last year,
    each area's count per build, postings by month, visits, and one row per quarter."""
    year_ago = (TODAY - dt.timedelta(days=365)).isoformat()
    recent = [h for h in hist if h["date"] >= year_ago]
    series = [{k: h.get(k) for k in ("date", "open", "interns", "employers", "new7", "new14", "medHr", "payListed", "remote", "bilingual")} for h in recent]
    area_names = sorted({a for h in recent for a in h.get("areas", {})})
    area_series = {a: [h.get("areas", {}).get(a, 0) for h in recent] for a in area_names}
    posted = collections.Counter()   # a month's postings: the most any build counted for it
    for h in hist:
        for m, n in h.get("postedByMonth", {}).items(): posted[m] = max(posted[m], n)
    quarters = []
    for q in sorted({quarter_of(h["date"]) for h in hist}):
        hs = [h for h in hist if quarter_of(h["date"]) == q]
        meds = sorted(h["medHr"] for h in hs if h.get("medHr") is not None)
        areas = collections.Counter()
        for h in hs: areas.update(h.get("areas", {}))
        emps = collections.Counter()
        for h in hs:
            for name, n in h.get("topEmployers", []): emps[name] += n
        quarters.append({
            "q": q, "builds": len(hs), "from": hs[0]["date"], "to": hs[-1]["date"],
            "openAvg": round(sum(h["open"] for h in hs) / len(hs)), "openHigh": max(h["open"] for h in hs),
            "employersAvg": round(sum(h["employers"] for h in hs) / len(hs)),
            "medHr": round(quant(meds, .5), 2) if meds else None,
            "payListedPct": round(100 * sum(h["payListed"] for h in hs) / max(1, sum(h["open"] for h in hs))),
            "posted": sum(n for m, n in posted.items() if quarter_of(m + "-01") == q),
            "topAreas": [a for a, _ in areas.most_common(3)], "topEmployers": [a for a, _ in emps.most_common(3)],
        })
    for q in quarters:   # a quarter's visits, where the record has them
        qd = [v for d, v in (visits or {}).items() if quarter_of(d) == q["q"]]
        q["visits"] = sum(int(v.get("board", 0)) + int(v.get("insights", 0)) + int(v.get("other", 0)) for v in qd) if qd else None
    return {"series": series, "areaSeries": area_series, "postedByMonth": sorted(posted.items()), "visits": visits_series(visits),
            "quarters": quarters, "first": hist[0]["date"] if hist else None, "builds": len(hist)}

def render_insights(template_path, ins, date_label):
    tpl = open(template_path, encoding='utf-8').read()
    for tok in ('__INSIGHTS__', '__DATE__'):
        if tpl.count(tok) != 1: sys.exit(f"{template_path}: expected exactly one {tok} placeholder")
    return tpl.replace('__INSIGHTS__', json_for_html(ins)).replace('__DATE__', date_label)

def render_assistant(template_path, date_label):
    """The assistant page carries no data: its listings come from /api/assistant, which reads
    the built index.html. Only the 'updated' label is filled in."""
    tpl = open(template_path, encoding='utf-8').read()
    if tpl.count('__DATE__') != 1: sys.exit(f"{template_path}: expected exactly one __DATE__ placeholder")
    return tpl.replace('__DATE__', date_label)

def json_for_html(out):
    """JSON that is safe to inline inside a <script type="application/json"> block."""
    return json.dumps(out, ensure_ascii=False, separators=(",", ":")).replace("<", "\\u003c")

def sha256_b64(text):
    return "sha256-" + base64.b64encode(hashlib.sha256(text.encode("utf-8")).digest()).decode()

def security_headers(pages):
    """Build a Content-Security-Policy that fits each page exactly.

    `pages` maps a route to the HTML served there (the board at "/" and "/index.html", the
    insights page at "/insights" and "/insights.html").

    Both the app script and the stylesheet are inline, so they are allowed by hash
    rather than by 'unsafe-inline'. That means an injected script cannot run even if
    escaping somewhere is wrong. The hashes cover the exact bytes of the rendered page,
    so index.html and vercel.json must be regenerated together; hand-editing index.html
    will stop its script from running.
    """
    def csp_for(route, html):
        scripts = re.findall(r"<script(?![^>]*\btype=)[^>]*>(.*?)</script>", html, re.S)
        styles = re.findall(r"<style[^>]*>(.*?)</style>", html, re.S)
        if not scripts: sys.exit(f"{route}: no inline script found: refusing to write a policy that would break the page")
        return "; ".join([
            "default-src 'none'",
            "script-src " + " ".join(f"'{sha256_b64(x)}'" for x in scripts) + " 'self'",
            "style-src " + " ".join(f"'{sha256_b64(x)}'" for x in styles) + " https://fonts.googleapis.com",
            "font-src https://fonts.gstatic.com",
            "img-src 'self' data:",
            # Vercel Analytics posts page views to a same-origin /_vercel/insights/* path, so
            # 'self' is what actually matters here; the vitals host covers Speed Insights.
            "connect-src 'self' https://vitals.vercel-insights.com",
            "form-action 'none'",
            "frame-ancestors 'none'",
            "base-uri 'none'",
        "object-src 'none'",
        ])
    # The CSP is written for each page's inline script and style, so it goes on the pages
    # only; a PDF or image served beside them (the guide deck) opens in the browser's own
    # viewer, which the policy would otherwise get in the way of. The rest apply everywhere.
    return {"headers": [
        {"source": route, "headers": [{"key": "Content-Security-Policy", "value": csp_for(route, html)}]} for route, html in pages.items()
    ] + [
        # The assistant's API answers JSON only; nothing on it may be framed or run as a page.
        {"source": "/api/(.*)", "headers": [
            {"key": "Content-Security-Policy", "value": "default-src 'none'; frame-ancestors 'none'"},
            {"key": "Cache-Control", "value": "no-store"},
        ]},
        {"source": "/(.*)", "headers": [
            {"key": "X-Content-Type-Options", "value": "nosniff"},
            {"key": "Referrer-Policy", "value": "strict-origin-when-cross-origin"},
            {"key": "Permissions-Policy", "value": "camera=(), microphone=(), geolocation=(), payment=(), usb=()"},
            {"key": "X-Frame-Options", "value": "DENY"},
            {"key": "Strict-Transport-Security", "value": "max-age=31536000; includeSubDomains"},
        ]},
    ],
    # The assistant function reads the listings out of the built index.html at run time, so
    # that file has to be bundled next to it. See "The opportunity assistant" in the README.
    "functions": {"api/assistant.js": {"maxDuration": 60, "includeFiles": "index.html"}}}

def render_html(template_path, out, date_label):
    tpl = open(template_path, encoding='utf-8').read()
    for tok in ('__DATA__', '__DATE__', '__HUBS__'):
        if tpl.count(tok) != 1: sys.exit(f"{template_path}: expected exactly one {tok} placeholder")
    labels = [{"id": h["id"], "label": h["label"], "note": h["note"]} for h in HUBS]
    return tpl.replace('__DATA__', json_for_html(out)).replace('__HUBS__', json_for_html(labels)).replace('__DATE__', date_label)

def print_stats(out, items):
    print(collections.Counter(x['cat'] for x in out).most_common())
    print(collections.Counter(x['level'] for x in out))
    print(collections.Counter(x['mode'] for x in out))
    print("pay listed", collections.Counter(bool(x['pay']) for x in out))
    print("flags", collections.Counter(x['flag'] for x in out))
    print("counties", collections.Counter(c for x in out for c in x['counties']))
    print("no majors", sum(1 for x in out if not x['majors']))
    print(collections.Counter(tuple(x['term']) for x in out).most_common(12))
    print(collections.Counter(m for x in out for m in x['majors']).most_common(50))

def main(argv=None):
    here = os.path.dirname(os.path.abspath(__file__))
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", nargs="?", help="spreadsheet export of the listings")
    ap.add_argument("--data", help="skip the spreadsheet and re-render from this data.json (or a built index.html)")
    ap.add_argument("--sheet", nargs=2, action="append", default=[], metavar=("AREA", "PATH"),
                    help="a hand-made .xlsx or .csv of jobs and internships for one spotlight area (one of: "
                         + "; ".join(SPOT_AREAS) + "). Columns: Company, Role, Location, Direct Apply URL (or the "
                         "names in SHEET_ALIASES), and optionally Posted, Type, Experience, Salary, Summary, Caveats. "
                         "Every row lands in that area. Repeatable; can stand in for the export, or be added to one.")
    ap.add_argument("--relink", metavar="XLSX",
                    help="an Indeed export: any card that applies through Indeed and whose posting is in it "
                         "is pointed at the employer's own apply link instead, keeping the Indeed page as its "
                         "original listing. Adds no cards.")
    ap.add_argument("--merge", metavar="PATH",
                    help="add this build to the listings already published, instead of replacing them. "
                         "PATH is a data.json or a built index.html; a role already on the board is "
                         "updated in place rather than added a second time.")
    ap.add_argument("--json", default=os.path.join(here, "data.json"), help="where to write data.json (default: next to this script)")
    ap.add_argument("--template", default=os.path.join(here, "template.html"))
    ap.add_argument("--out", default=os.path.join(here, "index.html"), help="rendered page (default: index.html next to this script); '-' to skip")
    ap.add_argument("--insights-template", default=os.path.join(here, "insights_template.html"))
    ap.add_argument("--history", default=os.path.join(here, "history.json"), help="where each build records its figures for the trends on the insights page (default: history.json next to this script); '-' to skip")
    ap.add_argument("--visits", default=os.path.join(here, "visits.json"), help="page views recorded by record_visits.py, for the insights page ('-' to skip)")
    ap.add_argument("--no-snapshot", action="store_true", help="redraw the pages from the history as it is, without adding today's snapshot (the weekly visits job)")
    ap.add_argument("--insights-out", default=os.path.join(here, "insights.html"), help="the insights page (default: insights.html next to this script); '-' to skip")
    ap.add_argument("--assistant-template", default=os.path.join(here, "assistant_template.html"))
    ap.add_argument("--assistant-out", default=os.path.join(here, "assistant.html"), help="the assistant page (default: assistant.html next to this script); '-' to skip")
    ap.add_argument("--date", default=f"{TODAY:%b} {TODAY.day}, {TODAY.year}", help="'updated' label, e.g. 'Sep 10, 2026'")
    ap.add_argument("--max-age", type=int, default=90, metavar="DAYS",
                    help="drop postings older than DAYS (default 90; use 0 to keep everything). "
                         "A dead apply link costs more trust than a missing listing earns.")
    ap.add_argument("--keep-http", action="store_true", help="leave plain-HTTP apply links alone instead of upgrading them to HTTPS")
    ap.add_argument("--strict", action="store_true", help="exit non-zero if a listing contains text shaped like an AI instruction")
    ap.add_argument("--headers", default=None, help="where to write vercel.json with the CSP and security headers ('-' to skip)")
    ap.add_argument("--featured", default=os.path.join(here, "featured.json"),
                    help="paid placements to pin at the top of the board (default: featured.json next to this script)")
    ap.add_argument("--stats", action="store_true", help="print classification counters")
    ap.add_argument("--check", action="store_true", help="also write check.txt, one line per card, for eyeballing categories")
    a = ap.parse_args(argv)
    for area, _ in a.sheet:
        if area not in SPOT_AREAS: ap.error(f"--sheet: unknown area {area!r}; one of: " + "; ".join(SPOT_AREAS))
    if not a.xlsx and not a.data and not a.sheet: ap.error("give a spreadsheet path, --sheet AREA PATH or --data data.json")
    cutoff = (TODAY - dt.timedelta(days=a.max_age)).strftime("%Y-%m-%d") if a.max_age else ""

    if a.data:
        out = load_cards(a.data)
        items, skipped = out, 0
        print(f"loaded {len(out)} cards from {a.data}")
        if a.max_age:
            out, stale = drop_stale(out, cutoff)
            if stale:
                print(f"dropped {len(stale)} card(s) last posted before {cutoff}: "
                      + "; ".join(f"{g['company']} / {g['title'][:40]} ({g['posted']})" for g in stale[:12])
                      + (" ..." if len(stale) > 12 else ""))
    else:
        data, items, skipped = [], [], 0
        if a.xlsx:
            data = read_rows(a.xlsx)
            items, skipped = build_items(data, upgrade=not a.keep_http)
        for area, path in a.sheet:
            cdata = read_sheet_rows(path)
            citems, cskipped = build_sheet_items(cdata, area, upgrade=not a.keep_http)
            print(f"read {len(cdata)} {area} row(s) from {path}: {len(citems)} usable, {cskipped} skipped "
                  f"({sum(1 for it in citems if it['kind'] == 'Job')} jobs, {sum(1 for it in citems if it['kind'] == 'Internship')} internships)")
            data += cdata; items += citems; skipped += cskipped
        if a.max_age:
            stale = [it for it in items if it['posted'] and it['posted'] < cutoff]
            items = [it for it in items if not (it['posted'] and it['posted'] < cutoff)]
            print(f"dropped {len(stale)} postings older than {cutoff}" + (": " + "; ".join(f"{it['company']} / {it['title'][:40]} ({it['posted']})" for it in stale[:8]) + (" ..." if len(stale) > 8 else "") if stale else ""))
        out = group_items(items)
        print(f"read {len(data)} rows, skipped {skipped} (blank or duplicate id), grouped {len(items)} postings into {len(out)} cards")
        if NOT_INTERNSHIPS:
            print(f"left out {len(NOT_INTERNSHIPS)} posting(s) that are jobs, not internships: " + "; ".join(NOT_INTERNSHIPS[:8]) + (" ..." if len(NOT_INTERNSHIPS) > 8 else ""))
    out = finalize(out, upgrade=not a.keep_http)
    if a.merge:
        prior = finalize(load_cards(a.merge), upgrade=not a.keep_http)
        fresh = len(out)
        out, added, updated, same = merge_cards(prior, out)
        print(f"merged into {len(prior)} card(s) already on the board from {a.merge}: "
              f"{len(added)} new, {len(updated)} updated, {len(same)} already there (of {fresh} built)")
        if added:
            print("  new: " + "; ".join(f"{g['company']} / {display_title(g['title'])[:40]}" for g in added[:8])
                  + (f" ... and {len(added)-8} more" if len(added) > 8 else ""))
        confirmed = {g.get('id') for g in added + updated + same}
        out = finalize(out, upgrade=not a.keep_http)
        out, closing = drop_closing(out, confirmed)
        if closing:
            print(f"dropped {len(closing)} card(s) for a term in its last two months that this export no longer lists: "
                  + "; ".join(f"{g['company']} / {g['title'][:40]} ({', '.join(g['term'])})" for g in closing[:12])
                  + (" ..." if len(closing) > 12 else ""))
        if a.max_age:
            out, stale = drop_stale(out, cutoff)
            if stale:
                print(f"dropped {len(stale)} published card(s) last posted before {cutoff}: "
                      + "; ".join(f"{g['company']} / {g['title'][:40]} ({g['posted']})" for g in stale[:8])
                      + (" ..." if len(stale) > 8 else ""))
    if UPGRADED_HTTP:
        print(f"upgraded {sum(UPGRADED_HTTP.values())} plain-HTTP link(s) to HTTPS across "
              + ", ".join(f"{d} x{n}" for d, n in UPGRADED_HTTP.most_common(6))
              + (" ..." if len(UPGRADED_HTTP) > 6 else "")
              + " (spot-check any unfamiliar domain, or rerun with --keep-http)")
    if a.relink:
        changed = relink_cards(out, read_rows(a.relink), upgrade=not a.keep_http)
        print(f"relinked {len(changed)} card(s) from Indeed to the employer's apply page using {a.relink}"
              + (": " + "; ".join(f"{g['company']} / {g['title'][:40]} -> {hostOf(g['locs'][0]['apply'])}" for g in changed) if changed else ""))
        out = finalize(out, upgrade=not a.keep_http)
    out, held = drop_excluded(out)
    if held:
        print(f"held back {len(held)} posting(s) listed in EXCLUDED: "
              + "; ".join(f"{g['company']} / {g['title'][:40]} ({why})" for g, why in held))
    out, expired = drop_expired(out)
    if expired:
        print(f"dropped {len(expired)} card(s) whose only term has already ended: "
              + "; ".join(f"{g['company']} / {g['title'][:40]} ({', '.join(g['term'])})" for g in expired[:8])
              + (" ..." if len(expired) > 8 else ""))
    out, folded = dedupe_cards(out)
    if folded:
        print(f"folded {len(folded)} duplicate card(s) into a newer copy of the same role: "
              + "; ".join(f"{g['company']} / {g['title'][:40]} ({g['posted']})" for g in folded[:8])
              + (" ..." if len(folded) > 8 else ""))
    feats, fadded, fmissing, fended = apply_featured(out, load_featured(a.featured), upgrade=not a.keep_http)
    if feats or fmissing or fended:
        print(f"featured {len(feats)} card(s) from {a.featured}"
              + (": " + "; ".join(f"{g['company']} / {g['title'][:40]} until {g['featured']}" for g in feats) if feats else ""))
        if fadded: print(f"  added {len(fadded)} featured role(s) no export carried: " + "; ".join(f"{g['company']} / {g['title'][:40]}" for g in fadded))
        if fended: print(f"  {len(fended)} placement(s) already ended, take them out of {a.featured}: " + "; ".join(f"{e['company']} / {e['title'][:40]} ({e['until']})" for e in fended))
        if fmissing: print(f"  WARNING: {len(fmissing)} featured role(s) are not on the board and give no apply link to add them with: " + "; ".join(f"{e['company']} / {e['title'][:40]}" for e in fmissing))
        out = finalize(out, upgrade=not a.keep_http)
    hits = scan_injection(out)
    if hits:
        print(f"WARNING: {len(hits)} listing field(s) contain text shaped like an AI instruction:")
        for company, title, field, frag in hits[:10]:
            print(f"  {company} / {title[:40]} [{field}]: {frag[:70]!r}")
        print("  These do not affect the site, which runs no model. Review before committing,")
        print("  because coding agents read this repo.")
        if a.strict:
            sys.exit(f"--strict: refusing to write anything while {len(hits)} field(s) look suspicious")
    else:
        print(f"injection scan: clean across {len(out)} listings")

    if a.stats: print_stats(out, items)

    if not a.data or os.path.abspath(a.data) != os.path.abspath(a.json):
        with open(a.json, 'w', encoding='utf-8') as f: json.dump(out, f, ensure_ascii=False)
        print(f"wrote {a.json} ({os.path.getsize(a.json):,} bytes)")
    if a.check:
        with open(os.path.join(os.path.dirname(os.path.abspath(a.json)), 'check.txt'), 'w', encoding='utf-8') as f:
            for x in sorted(out, key=lambda x: x['cat']):
                f.write(f"{x['cat'][:30]:30} | {x['company'][:22]:22} | {x['title'][:55]:55} | {x['level'][:8]:8} | {x['pay'][:14]:14} | {x['flag'][:10]:10} | {x['count']} | {','.join(x['majors'])[:60]}\n")
    if a.out != '-':
        html = render_html(a.template, out, a.date)
        with open(a.out, 'w', encoding='utf-8') as f: f.write(html)
        print(f"wrote {a.out} ({len(html):,} chars, updated {a.date})")
        pages = {"/": html, "/index.html": html}
        if a.insights_out != '-':
            ins = insights_for(out)
            if a.history != '-':
                if a.no_snapshot:
                    hist = json.load(open(a.history, encoding='utf-8')) if os.path.exists(a.history) else []
                else:
                    hist = update_history(a.history, snapshot_for(out, ins))
                    print(f"wrote {a.history} ({len(hist)} snapshot(s), {hist[0]['date']} to {hist[-1]['date']})")
                visits = load_visits(a.visits)
                ins["history"] = history_series(hist, visits)
                if visits: print(f"read {len(visits)} day(s) of visits from {a.visits}")
            ihtml = render_insights(a.insights_template, ins, a.date)
            with open(a.insights_out, 'w', encoding='utf-8') as f: f.write(ihtml)
            print(f"wrote {a.insights_out} ({len(ihtml):,} chars)")
            pages.update({"/insights": ihtml, "/insights.html": ihtml})
        if a.assistant_out != '-':
            ahtml = render_assistant(a.assistant_template, a.date)
            with open(a.assistant_out, 'w', encoding='utf-8') as f: f.write(ahtml)
            print(f"wrote {a.assistant_out} ({len(ahtml):,} chars)")
            pages.update({"/assistant": ahtml, "/assistant.html": ahtml})
        headers_path = a.headers or os.path.join(os.path.dirname(os.path.abspath(a.out)), 'vercel.json')
        if headers_path != '-':
            with open(headers_path, 'w', encoding='utf-8') as f:
                json.dump(security_headers(pages), f, indent=2); f.write("\n")
            print(f"wrote {headers_path} (CSP pinned to each page's inline script and style)")


if __name__ == "__main__":
    main()
